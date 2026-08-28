"""Excel report exports (openpyxl workbooks streamed as .xlsx downloads)."""
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lib.db import db

router = APIRouter()

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)


def _write_sheet(wb: Workbook, title: str, heading: str, headers: List[str], rows: List[List[Any]]) -> None:
    ws = wb.active if wb.sheetnames == ["Sheet"] else wb.create_sheet()
    ws.title = title

    ws.append([heading])
    ws["A1"].font = TITLE_FONT
    ws.append([f"Dicetak: {datetime.now(timezone.utc).strftime('%d-%m-%Y %H:%M')} UTC"])
    ws.append([])

    ws.append(headers)
    header_row = ws.max_row
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append(row)

    widths = [len(str(h)) for h in headers]
    for row in rows:
        for i, value in enumerate(row):
            widths[i] = max(widths[i], min(len(str(value)), 48))
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width + 4

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def _stream(wb: Workbook, filename: str) -> StreamingResponse:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type=XLSX_MIME,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/reports/products.xlsx")
async def export_products():
    docs = await db.products.find().sort("name", 1).to_list(2000)
    rows: List[List[Any]] = []
    for p in docs:
        stock = int(p.get("current_stock", 0))
        damaged = int(p.get("damaged_stock", 0))
        buy = float(p.get("purchase_price", 0))
        min_stock = int(p.get("min_stock", 0))
        ups = int(p.get("units_per_secondary", 1)) or 1
        rows.append([
            p.get("name", ""), p.get("sku", ""), p.get("category", ""), p.get("unit", ""),
            stock, damaged, p.get("expiry_date") or "", min_stock,
            "PERLU RESTOCK" if min_stock > 0 and stock <= min_stock else "AMAN",
            f"{p.get('weight_per_unit', 1)} {p.get('weight_unit', 'Kg')}",
            f"{ups} {p.get('unit', 'Pcs')}/{p.get('secondary_unit', 'Dus')}",
            round(stock * float(p.get("weight_per_unit", 1)), 3),
            round(stock / ups, 2),
            buy, float(p.get("selling_price", 0)), buy * stock,
            p.get("supplier_name", ""), p.get("location", ""),
        ])
    total_value = sum(r[15] for r in rows)
    rows.append(["TOTAL", "", "", "",
                 sum(int(r[4]) for r in rows), sum(int(r[5]) for r in rows),
                 "", "", "", "", "", "", "", "", "", total_value, "", ""])

    wb = Workbook()
    _write_sheet(
        wb, "Daftar Produk", "Laporan Daftar Produk & Nilai Stok",
        ["Nama Produk", "SKU", "Kategori", "Satuan", "Stok Baik", "Stok Rusak", "Tanggal EXP",
         "Stok Minimum", "Status Stok",
         "Berat/Satuan", "Isi Kemasan Sekunder", "Total Berat", "Total Kemasan Sekunder",
         "Harga Modal (Rp)", "Harga Jual (Rp)", "Nilai Total (Rp)", "Supplier", "Lokasi"],
        rows,
    )
    return _stream(wb, f"laporan-produk-{datetime.now(timezone.utc).date().isoformat()}.xlsx")


@router.get("/reports/stock-locations.xlsx")
async def export_stock_locations():
    """Rekap stok per tumpukan + perkalian tumpukan P × L × T."""
    docs = await db.placements.find().to_list(5000)

    def sort_key(d: Dict[str, Any]) -> Any:
        unit = str(d.get("unit_name", ""))
        digits = "".join(ch for ch in unit if ch.isdigit())
        return ("".join(ch for ch in unit if ch.isalpha()), int(digits or 0), str(d.get("stack", "")))

    docs.sort(key=sort_key)

    rows: List[List[Any]] = []
    for p in docs:
        rows.append([
            p.get("complex_name", ""), p.get("unit_name", ""), p.get("stack", ""),
            p.get("location_code", ""),
            p.get("product_name", ""), p.get("product_sku", ""),
            f"{p.get('units_per_secondary', 1)} {p.get('unit', 'Pcs')}/{p.get('secondary_unit', 'Karung')}",
            f"{p.get('weight_per_secondary', 0)} {p.get('weight_unit', 'Kg')}",
            int(p.get("length", 1)), int(p.get("width", 1)), int(p.get("height", 1)),
            int(p.get("secondary_count", 0)),
            int(p.get("total_units", 0)),
            round(float(p.get("total_weight", 0)), 3),
            p.get("created_by_name", ""), p.get("notes", ""),
        ])

    rows.append([])
    rows.append([
        "TOTAL", "", "", "", "", "", "", "", "", "", "",
        sum(int(r[11]) for r in rows if len(r) > 11 and isinstance(r[11], int)),
        sum(int(r[12]) for r in rows if len(r) > 12 and isinstance(r[12], int)),
        round(sum(float(r[13]) for r in rows if len(r) > 13 and isinstance(r[13], (int, float))), 3),
        "", "",
    ])

    wb = Workbook()
    _write_sheet(
        wb, "Stok per Tumpukan", "Rekap Stok per Tumpukan — Bulog Gudang Sunter Timur I & II",
        ["Kompleks Gudang", "Unit Gudang", "Tumpukan", "Kode Tumpukan", "Komoditas", "SKU",
         "Isi Kemasan Sekunder", "Berat per Kemasan", "P (Panjang)", "L (Lebar)", "T (Tinggi)",
         "Perkalian Tumpukan (P×L×T)", "Total Satuan Primer", "Total Berat (Kg)",
         "Dicatat Oleh", "Catatan"],
        rows,
    )
    return _stream(wb, f"rekap-tumpukan-{datetime.now(timezone.utc).date().isoformat()}.xlsx")


@router.get("/reports/transactions.xlsx")
async def export_transactions(month: Optional[str] = Query(None, description="Filter bulan YYYY-MM")):
    query: Dict[str, Any] = {}
    if month:
        query["date"] = {"$regex": f"^{month}"}
    docs = await db.transactions.find(query).sort("created_at", -1).to_list(5000)

    rows: List[List[Any]] = []
    for t in docs:
        created = t.get("created_at")
        stamp = created.strftime("%d-%m-%Y %H:%M") if isinstance(created, datetime) else t.get("date", "")
        rows.append([
            stamp, t.get("reference_no", ""), t.get("queue_no", ""), t.get("type", ""),
            t.get("product_name", ""), t.get("product_sku", ""), t.get("category", ""),
            int(t.get("quantity", 0)), t.get("condition", "BAIK"), int(t.get("stock_after", 0)),
            t.get("party", ""), t.get("vehicle_plate", ""),
            t.get("created_by_name", ""), t.get("notes", ""),
        ])

    masuk = sum(r[7] for r in rows if r[3] == "MASUK")
    keluar = sum(r[7] for r in rows if r[3] == "KELUAR")
    rows.append([])
    rows.append(["RINGKASAN", "", "", "MASUK", "", "", "", masuk, "", "", "", "", "", ""])
    rows.append(["RINGKASAN", "", "", "KELUAR", "", "", "", keluar, "", "", "", "", "", ""])

    label = f"Bulan {month}" if month else "Semua Periode"
    wb = Workbook()
    _write_sheet(
        wb, "Riwayat Transaksi", f"Laporan Riwayat Transaksi Stok — {label}",
        ["Waktu", "No. Referensi", "No. Antrian", "Tipe", "Nama Produk", "SKU", "Kategori",
         "Jumlah", "Kondisi Barang", "Stok Akhir", "Pihak Terkait", "No. Polisi Kendaraan",
         "Dicatat Oleh", "Catatan"],
        rows,
    )
    suffix = month or "semua"
    return _stream(wb, f"laporan-transaksi-{suffix}.xlsx")
