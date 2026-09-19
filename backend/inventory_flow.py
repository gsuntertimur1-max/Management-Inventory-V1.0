import csv
import io
from openpyxl import load_workbook
import random
import re
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from typing import List, Literal

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from backend.server import (
    build_xlsx,
    db,
    get_current_user,
    max_suffix,
    new_id,
    next_sequence,
    now_iso,
    operational_now,
    require_write,
    require_master_write,
    ensure_channel_stock,
    normalize_channel,
    get_operational_location,
)
from backend.stack_allocations import allocate_stock_to_stack, decrease_stack_allocation
from backend.work_time_costs import handling_fee, holiday_from_settings, work_split

router = APIRouter(prefix="/api")


class POItemInput(BaseModel):
    productId: str
    qty: float = Field(gt=0)


class PurchaseOrderInput(BaseModel):
    supplier: str
    no: str = ""
    items: List[POItemInput] = Field(min_length=1)
    date: str = ""


class PurchaseOrderCancelInput(BaseModel):
    reason: str = Field(min_length=3, max_length=500)


class ReceiptItemInput(BaseModel):
    productId: str
    # qty/kondisi is retained for older clients. New receipt forms send the
    # two quantities in a single receiving transaction.
    qty: float = Field(default=0, ge=0)
    goodQty: float = Field(default=0, ge=0)
    damagedQty: float = Field(default=0, ge=0)
    normalQtyBefore1600: float | None = Field(default=None, ge=0)
    exp: str = ""
    stackCode: str = ""
    channel: str = ""


class ReceiptInput(BaseModel):
    poId: str = ""
    items: List[ReceiptItemInput] = Field(min_length=1)
    party: str = ""
    ref: str = ""
    polisi: str = ""
    kondisi: Literal["BAIK", "RUSAK"] = "BAIK"
    keterangan: str = ""
    weighingForm: bool = False
    grossWeight: float = Field(default=0, ge=0)
    grossMin: float = Field(default=0, ge=0)
    grossMax: float = Field(default=0, ge=0)
    # PENGIRIM ditagihkan pada pengirim; TERMAKSUK berarti sudah masuk harga/dokumen.
    unloadingFeeChargeMode: Literal["", "PENGIRIM", "TERMASUK"] = ""
    unloadingSessionId: str = ""
    unloadingStartTime: str = ""  # legacy client compatibility only


def _unloading_user_key(user: dict) -> str:
    return str(user.get("id") or user.get("username") or user.get("name") or "").strip()


@router.post("/unloading-sessions/start")
async def start_unloading_session(user: dict = Depends(require_write)):
    user_key = _unloading_user_key(user)
    now = operational_now()
    existing = await db.unloading_sessions.find_one(
        {"startedByKey": user_key, "status": "BERJALAN"},
        {"_id": 0},
    )
    if existing:
        try:
            started = datetime.fromisoformat(str(existing.get("startedAt") or "").replace("Z", "+00:00"))
            if started.tzinfo is None:
                started = started.replace(tzinfo=now.tzinfo)
            else:
                started = started.astimezone(now.tzinfo)
        except ValueError:
            started = None
        if started and started.date() == now.date():
            return existing
        await db.unloading_sessions.update_one(
            {"id": existing.get("id"), "status": "BERJALAN"},
            {"$set": {"status": "DIBATALKAN_OTOMATIS", "cancelledAt": now_iso(), "cancelledBy": "Sistem"}},
        )
    doc = {
        "id": new_id(),
        "status": "BERJALAN",
        "startedAt": now.isoformat(),
        "startedBy": user.get("name", ""),
        "startedByKey": user_key,
        "createdAt": now_iso(),
    }
    await db.unloading_sessions.insert_one(dict(doc))
    return doc


@router.post("/unloading-sessions/{session_id}/cancel")
async def cancel_unloading_session(session_id: str, user: dict = Depends(require_write)):
    user_key = _unloading_user_key(user)
    result = await db.unloading_sessions.update_one(
        {"id": session_id, "startedByKey": user_key, "status": "BERJALAN"},
        {"$set": {"status": "DIBATALKAN", "cancelledAt": now_iso(), "cancelledBy": user.get("name", "")}},
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=409, detail="Sesi bongkar tidak aktif atau bukan milik pengguna ini")
    return {"ok": True}


def receipt_condition_quantities(item: ReceiptItemInput, kondisi: str) -> tuple[float, float]:
    good_qty = float(item.goodQty or 0)
    damaged_qty = float(item.damagedQty or 0)
    if good_qty <= 0 and damaged_qty <= 0:
        if float(item.qty or 0) <= 0:
            raise HTTPException(status_code=400, detail="Jumlah baik atau rusak harus diisi")
        return (0.0, float(item.qty)) if kondisi == "RUSAK" else (float(item.qty), 0.0)
    total_qty = good_qty + damaged_qty
    if float(item.qty or 0) > 0 and abs(float(item.qty) - total_qty) > 1e-9:
        raise HTTPException(status_code=400, detail="Total penerimaan harus sama dengan jumlah baik dan rusak")
    return good_qty, damaged_qty


class DamageDiscoveryInput(BaseModel):
    productId: str
    stackCode: str
    qty: float = Field(gt=0)
    channel: str = ""
    cause: str = Field(min_length=3, max_length=200)
    note: str = ""
    referenceNo: str = ""


class SupplierReturnInput(BaseModel):
    productId: str
    qty: float = Field(gt=0)
    channel: str = ""
    supplier: str = ""
    sourceDamageOperationId: str = ""
    sourceStackCode: str = ""
    poNo: str = ""
    returnNo: str = ""
    note: str = ""


class SupplierReplacementInput(BaseModel):
    qty: float = Field(gt=0)
    stackCode: str
    referenceNo: str = ""
    note: str = ""


def _number(value, default=0.0) -> float:
    if value is None or value == "":
        return default
    text = str(value).strip().replace(" ", "")
    if not text:
        return default
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


def _unloading_fee(product: dict, qty: float, at, charge_mode_override: str = "", overtime_qty: float | None = None, holiday_override: bool | None = None) -> dict:
    if overtime_qty is None:
        overtime_qty = float(qty or 0) if at.hour >= 16 else 0.0
    return handling_fee(product, qty, overtime_qty, at, "unloading", "PENGIRIM", charge_mode_override, holiday_override)


def _unloading_started_at(value: str, completed_at: datetime) -> datetime:
    text = str(value or "").strip()
    if not text:
        return completed_at
    try:
        hour_text, minute_text = text.split(":", 1)
        hour, minute = int(hour_text), int(minute_text)
        if not (0 <= hour <= 23 and 0 <= minute <= 59):
            raise ValueError
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Waktu mulai bongkar harus HH:MM") from exc
    started = completed_at.replace(hour=hour, minute=minute, second=0, microsecond=0)
    if started > completed_at:
        started -= timedelta(days=1)
    return started


def _validate_exp(value: str) -> str:
    value = (value or "").strip()
    if not value:
        return ""
    try:
        datetime.strptime(value, "%Y-%m-%d")
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Format tanggal kedaluwarsa harus YYYY-MM-DD") from exc
    return value


def _validate_pack_qty(product: dict, qty: float) -> None:
    if float(product.get("secondaryQty", 0) or 0) > 0 and abs(qty - round(qty)) > 1e-6:
        raise HTTPException(
            status_code=400,
            detail=f"Jumlah {product.get('name', 'produk')} harus berupa kemasan primer/pack utuh",
        )


def _weighing_entries(average: float, minimum: float, maximum: float) -> list[dict]:
    target = round(average * 100)
    low, high = round(minimum * 100), round(maximum * 100)
    spread = min(target - low, high - target)
    chooser = random.SystemRandom()
    values = []
    for _ in range(10):
        delta = chooser.randint(0, spread)
        values.extend([target - delta, target + delta])
    chooser.shuffle(values)
    return [{"no": index, "gross": value / 100} for index, value in enumerate(values, 1)]


def _po_status(items: list[dict]) -> str:
    ordered = sum(float(item.get("qty", 0) or 0) for item in items)
    received = sum(float(item.get("receivedQty", item.get("received_qty", 0)) or 0) for item in items)
    cancelled = sum(float(item.get("cancelledQty", 0) or 0) for item in items)
    if cancelled > 1e-9:
        return "Dibatalkan" if received <= 1e-9 else "Diterima Sebagian · Sisa Dibatalkan"
    if ordered <= 0 or received <= 0:
        return "Belum Diterima"
    if received + 1e-9 < ordered:
        return "Sebagian"
    return "Selesai"


def _normalize_po(doc: dict) -> dict:
    doc = dict(doc)
    doc.pop("_id", None)
    normalized_items = []
    for item in doc.get("items", []):
        normalized = dict(item)
        normalized["receivedQty"] = float(
            normalized.get("receivedQty", normalized.get("received_qty", 0)) or 0
        )
        normalized.setdefault("unit", "")
        normalized.setdefault("sku", "")
        normalized.setdefault("productId", "")
        normalized_items.append(normalized)
    doc["items"] = normalized_items
    doc["status"] = _po_status(normalized_items)
    return doc


async def _hydrate_legacy_po(doc: dict) -> dict:
    """Lengkapi PO lama yang belum menyimpan productId/satuan agar tetap bisa dipakai."""
    normalized = _normalize_po(doc)
    changed = False
    for item in normalized.get("items", []):
        if item.get("productId"):
            continue
        query = {"sku": item.get("sku")} if item.get("sku") else {"name": item.get("name", "")}
        product = await db.products.find_one(query, {"_id": 0}) if query else None
        if not product:
            continue
        item["productId"] = product.get("id", "")
        item["sku"] = product.get("sku", "")
        item["unit"] = product.get("unit", "")
        item["cost"] = float(item.get("cost", product.get("cost", 0)) or 0)
        changed = True
    normalized["status"] = _po_status(normalized["items"])
    if changed and normalized.get("id"):
        await db.purchase_orders.update_one(
            {"id": normalized["id"]},
            {"$set": {"items": normalized["items"], "status": normalized["status"]}},
        )
    return normalized


@router.get("/purchase-orders-v2")
async def list_purchase_orders(user: dict = Depends(get_current_user)):
    docs = await db.purchase_orders.find({}, {"_id": 0}).sort("date", -1).to_list(1000)
    return [await _hydrate_legacy_po(doc) for doc in docs]


@router.post("/purchase-orders-v2")
async def create_purchase_order(body: PurchaseOrderInput, user: dict = Depends(require_master_write)):
    supplier = body.supplier.strip()
    if not supplier:
        raise HTTPException(status_code=400, detail="Supplier wajib dipilih")

    if not await db.suppliers.find_one({"name": supplier}):
        raise HTTPException(status_code=404, detail="Supplier tidak ditemukan")

    seen = set()
    items = []
    for item in body.items:
        if item.productId in seen:
            raise HTTPException(status_code=400, detail="Produk yang sama tidak boleh muncul dua kali dalam satu PO")
        seen.add(item.productId)

        product = await db.products.find_one({"id": item.productId}, {"_id": 0})
        if not product:
            raise HTTPException(status_code=404, detail="Produk pada PO tidak ditemukan")
        _validate_pack_qty(product, float(item.qty))

        items.append({
            "productId": product["id"],
            "sku": product.get("sku", ""),
            "name": product.get("name", ""),
            "qty": float(item.qty),
            "receivedQty": 0.0,
            "unit": product.get("unit", ""),
            "weight": float(product.get("weight", 0) or 0),
            "secondary": product.get("secondary", ""),
            "secondaryQty": float(product.get("secondaryQty", 0) or 0),
            "cost": float(product.get("cost", 0) or 0),
            "channel": normalize_channel(product.get("channel")),
        })

    manual_no = body.no.strip()
    if len(manual_no) > 100:
        raise HTTPException(status_code=400, detail="Nomor PO maksimal 100 karakter")
    if manual_no and await db.purchase_orders.find_one({"no": manual_no}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="Nomor PO sudah digunakan")
    if manual_no:
        po_no = manual_no
    else:
        year = operational_now().strftime("%Y")
        prefix = f"PO-{year}-"
        floor = await max_suffix(db.purchase_orders, "no", prefix)
        number = await next_sequence(f"purchase-order:{year}", floor)
        po_no = f"{prefix}{number:03d}"

    doc = {
        "id": new_id(),
        "no": po_no,
        "supplier": supplier,
        "date": body.date or now_iso(),
        "status": "Belum Diterima",
        "items": items,
        "total": sum(item["qty"] * item["cost"] for item in items),
        "created_by": user.get("name", ""),
        "created_at": now_iso(),
    }
    await db.purchase_orders.insert_one(dict(doc))
    return doc


@router.post("/purchase-orders-v2/{po_id}/cancel")
async def cancel_purchase_order(po_id: str, body: PurchaseOrderCancelInput, user: dict = Depends(require_master_write)):
    """Batalkan sisa PO yang belum diterima, tanpa mengubah penerimaan yang sudah tercatat."""
    raw_po = await db.purchase_orders.find_one({"id": po_id}, {"_id": 0})
    if not raw_po:
        raise HTTPException(status_code=404, detail="Purchase Order tidak ditemukan")
    po = await _hydrate_legacy_po(raw_po)
    if po.get("status") == "Selesai":
        raise HTTPException(status_code=400, detail="PO sudah selesai diterima dan tidak dapat dibatalkan")
    if "Dibatalkan" in str(po.get("status", "")):
        raise HTTPException(status_code=400, detail="Sisa PO sudah dibatalkan")

    items = []
    cancelled_total = 0.0
    for item in po.get("items", []):
        row = dict(item)
        remaining = max(float(row.get("qty", 0) or 0) - float(row.get("receivedQty", 0) or 0), 0)
        row["cancelledQty"] = remaining
        cancelled_total += remaining
        items.append(row)
    if cancelled_total <= 1e-9:
        raise HTTPException(status_code=400, detail="Tidak ada sisa PO yang dapat dibatalkan")

    event = {"id": new_id(), "reason": body.reason.strip(), "cancelledAt": now_iso(), "cancelledBy": user.get("name", ""), "qty": cancelled_total}
    updated = {**po, "items": items, "status": _po_status(items), "cancellation_history": list(po.get("cancellation_history") or []) + [event], "updated_at": now_iso()}
    await db.purchase_orders.update_one({"id": po_id}, {"$set": {key: value for key, value in updated.items() if key != "id"}})
    return updated


@router.post("/receipts")
async def receive_stock(body: ReceiptInput, user: dict = Depends(require_write)):
    if body.weighingForm and (body.grossWeight <= 0 or body.grossMin <= 0 or body.grossMax <= 0):
        raise HTTPException(status_code=400, detail="Rata-rata bruto serta rentang timbang harus diisi")
    if body.weighingForm and not body.grossMin <= body.grossWeight <= body.grossMax:
        raise HTTPException(status_code=400, detail="Rata-rata bruto harus berada di dalam rentang timbang")
    po = None
    if body.poId:
        raw_po = await db.purchase_orders.find_one({"id": body.poId}, {"_id": 0})
        if not raw_po:
            raise HTTPException(status_code=404, detail="Purchase Order tidak ditemukan")
        po = await _hydrate_legacy_po(raw_po)
        if po["status"] == "Selesai" or "Dibatalkan" in str(po.get("status", "")):
            raise HTTPException(status_code=400, detail="PO sudah selesai atau sisa penerimaannya telah dibatalkan")

    products = []
    requested_by_product = defaultdict(float)
    for item in body.items:
        product = await db.products.find_one({"id": item.productId}, {"_id": 0})
        if not product:
            raise HTTPException(status_code=404, detail="Produk penerimaan tidak ditemukan")
        good_qty, damaged_qty = receipt_condition_quantities(item, body.kondisi)
        total_qty = good_qty + damaged_qty
        if total_qty <= 0:
            raise HTTPException(status_code=400, detail="Jumlah baik atau rusak harus diisi")
        if good_qty > 0:
            _validate_pack_qty(product, good_qty)
        if damaged_qty > 0:
            _validate_pack_qty(product, damaged_qty)
        products.append({**product, "_channel": normalize_channel(item.channel, normalize_channel(product.get("channel"))), "_good_qty": good_qty, "_damaged_qty": damaged_qty})
        requested_by_product[item.productId] += total_qty
        _validate_exp(item.exp)

    party = (po.get("supplier") if po else body.party).strip() if (po or body.party) else ""
    if not party:
        raise HTTPException(status_code=400, detail="Supplier pengirim wajib dipilih")

    if po:
        po_item_map = {item.get("productId", ""): item for item in po.get("items", []) if item.get("productId")}
        for product_id, qty in requested_by_product.items():
            po_item = po_item_map.get(product_id)
            if not po_item:
                raise HTTPException(status_code=400, detail="Produk penerimaan tidak tercantum pada PO")
            ordered = float(po_item.get("qty", 0) or 0)
            received = float(po_item.get("receivedQty", 0) or 0)
            remaining = max(ordered - received, 0)
            if qty > remaining + 1e-9:
                raise HTTPException(
                    status_code=400,
                    detail=f"Jumlah diterima untuk {po_item.get('name', 'produk')} melebihi sisa PO ({remaining:g} {po_item.get('unit', '')})",
                )

    op_now = operational_now()
    unloading_session = None
    if body.unloadingSessionId.strip():
        unloading_session = await db.unloading_sessions.find_one(
            {
                "id": body.unloadingSessionId.strip(),
                "startedByKey": _unloading_user_key(user),
                "status": "BERJALAN",
            },
            {"_id": 0},
        )
        if not unloading_session:
            raise HTTPException(status_code=409, detail="Sesi bongkar tidak aktif. Tekan Mulai Bongkar kembali.")
        try:
            unloading_started_at = datetime.fromisoformat(str(unloading_session.get("startedAt") or "").replace("Z", "+00:00"))
            if unloading_started_at.tzinfo is None:
                unloading_started_at = unloading_started_at.replace(tzinfo=op_now.tzinfo)
            else:
                unloading_started_at = unloading_started_at.astimezone(op_now.tzinfo)
        except ValueError as exc:
            raise HTTPException(status_code=409, detail="Waktu sesi bongkar tidak valid. Batalkan sesi dan mulai kembali.") from exc
    else:
        unloading_started_at = _unloading_started_at(body.unloadingStartTime, op_now)
    fee_settings = await db.settings.find_one({"_id": "app"}, {"_id": 0, "holidays": 1}) or {}
    unloading_holiday = holiday_from_settings(op_now, fee_settings.get("holidays") or [])
    operation_id = new_id()
    transaction_ref = po.get("no") if po else (body.ref.strip() or f"IN-{op_now.strftime('%Y%m%d%H%M%S%f')}")
    time = now_iso()
    stock_changes = []
    stack_changes = []
    txns = []
    unloading_session_completed = False

    try:
        for item, product in zip(body.items, products):
            channel = product.get("_channel", normalize_channel(product.get("channel")))
            await ensure_channel_stock(product)
            exp = _validate_exp(item.exp)
            previous_exp = product.get("exp", "") or ""
            good_qty = float(product.get("_good_qty", 0) or 0)
            damaged_qty = float(product.get("_damaged_qty", 0) or 0)
            total_qty = good_qty + damaged_qty
            increments = {}
            if good_qty > 0:
                increments.update({"stock": good_qty, f"channelStock.{channel}.stock": good_qty})
            if damaged_qty > 0:
                increments.update({"damaged": damaged_qty, f"channelStock.{channel}.damaged": damaged_qty})
            update_doc = {"$inc": increments}

            # Di master produk hanya disimpan expired terdekat sebagai ringkasan.
            # Expired aktual tiap penerimaan tetap tersimpan pada transaksi.
            if good_qty > 0 and exp and (not previous_exp or exp < previous_exp):
                update_doc["$set"] = {"exp": exp}

            result = await db.products.update_one({"id": product["id"]}, update_doc)
            if result.matched_count == 0:
                raise HTTPException(status_code=409, detail=f"Produk {product.get('name', '')} berubah. Muat ulang lalu coba lagi.")

            for field, qty in (("stock", good_qty), ("damaged", damaged_qty)):
                if qty > 0:
                    stock_changes.append({"productId": product["id"], "field": field, "qty": qty, "previousExp": previous_exp, "expChanged": bool(update_doc.get("$set")), "channel": channel})

            stack_code = item.stackCode.strip().upper()
            location_source = stack_code or str(product.get("location", "") or "").strip()
            location_config = await get_operational_location(location_source, "inbound") if location_source else None
            if good_qty > 0 and not stack_code:
                raise HTTPException(status_code=400, detail=f"Pilih tumpukan penerimaan untuk {product.get('name', '')}")
            if good_qty > 0 and stack_code:
                from backend.stack_allocations import allocate_stock_to_stack
                await allocate_stock_to_stack(product, stack_code, good_qty, user.get("name", ""))
                stack_changes.append({"productId": product["id"], "stackCode": stack_code, "qty": good_qty})

            unloading_group = ""
            unloading_fee = {}
            unloading_split = work_split(
                unloading_started_at,
                op_now,
                total_qty,
                item.normalQtyBefore1600 if (body.unloadingSessionId or body.unloadingStartTime) else None,
            )
            if location_config and bool(location_config.get("unloadingCostEnabled", False)):
                unloading_group = str(location_config.get("unloadingGroup", "") or "")
                if unloading_group:
                    unloading_fee = _unloading_fee(
                        product,
                        total_qty,
                        op_now,
                        body.unloadingFeeChargeMode,
                        overtime_qty=unloading_split["overtimeQty"],
                        holiday_override=unloading_holiday,
                    )
                    unloading_fee.update({
                        **unloading_split,
                        "startedAt": unloading_started_at.isoformat(),
                        "completedAt": op_now.isoformat(),
                        "cutoff": "16:00",
                    })

            def receipt_txn(kondisi: str, qty: float, loading_cost: dict, include_weighing: bool):
                return {
                "id": new_id(),
                "operation_id": operation_id,
                "time": time,
                "operational_date": op_now.strftime("%Y-%m-%d"),
                "ref": transaction_ref,
                "po_id": po.get("id", "") if po else "",
                "po_no": po.get("no", "") if po else "",
                "antrian": "",
                "type": "MASUK",
                "kondisi": kondisi,
                "product": product.get("name", ""),
                "sku": product.get("sku", ""),
                "change": qty,
                "good_change": qty if kondisi == "BAIK" else 0,
                "damaged_change": qty if kondisi == "RUSAK" else 0,
                "unit": product.get("unit", ""),
                "weight": float(product.get("weight", 0) or 0),
                "total_weight": float(product.get("weight", 0) or 0) * qty,
                "secondary": product.get("secondary", ""),
                "secondaryQty": float(product.get("secondaryQty", 0) or 0),
                "channel": channel,
                "exp": exp,
                "penerima": party,
                "polisi": body.polisi,
                "operator": user.get("name", ""),
                "keterangan": body.keterangan,
                "weighing_form": body.weighingForm and include_weighing,
                "gross_weight": float(body.grossWeight) if body.weighingForm and include_weighing else 0,
                "gross_min": float(body.grossMin) if body.weighingForm and include_weighing else 0,
                "gross_max": float(body.grossMax) if body.weighingForm and include_weighing else 0,
                "weighing_entries": _weighing_entries(float(body.grossWeight), float(body.grossMin), float(body.grossMax)) if body.weighingForm and include_weighing else [],
                "unloading_group": unloading_group,
                "unloading_cost": loading_cost,
                "unloading_work": ({
                    "regularQty": loading_cost.get("regularQty", 0),
                    "overtimeQty": loading_cost.get("overtimeQty", 0),
                    "workStatus": loading_cost.get("workStatus", ""),
                    "startedAt": loading_cost.get("startedAt", ""),
                    "completedAt": loading_cost.get("completedAt", ""),
                    "cutoff": loading_cost.get("cutoff", "16:00"),
                } if loading_cost else {}),
            }

            if good_qty > 0:
                txns.append(receipt_txn("BAIK", good_qty, unloading_fee, True))
            if damaged_qty > 0:
                damaged_cost = unloading_fee if good_qty <= 0 else {}
                txns.append(receipt_txn("RUSAK", damaged_qty, damaged_cost, good_qty <= 0))

        if txns:
            await db.transactions.insert_many([dict(txn) for txn in txns])

        if unloading_session:
            session_result = await db.unloading_sessions.update_one(
                {"id": unloading_session["id"], "status": "BERJALAN"},
                {"$set": {
                    "status": "SELESAI",
                    "completedAt": op_now.isoformat(),
                    "completedBy": user.get("name", ""),
                    "operationId": operation_id,
                }},
            )
            if session_result.matched_count == 0:
                raise HTTPException(status_code=409, detail="Sesi bongkar sudah digunakan. Muat ulang halaman.")
            unloading_session_completed = True

        updated_po = None
        if po:
            updated_items = []
            for item in po["items"]:
                updated = dict(item)
                product_id = updated.get("productId", "")
                if product_id in requested_by_product:
                    updated["receivedQty"] = float(updated.get("receivedQty", 0) or 0) + requested_by_product[product_id]
                updated_items.append(updated)

            new_status = _po_status(updated_items)
            result = await db.purchase_orders.update_one(
                {"id": po["id"]},
                {"$set": {
                    "items": updated_items,
                    "status": new_status,
                    "last_received_at": time,
                    "last_received_by": user.get("name", ""),
                }},
            )
            if result.matched_count == 0:
                raise HTTPException(status_code=409, detail="PO berubah. Muat ulang lalu coba kembali.")
            updated_po = {**po, "items": updated_items, "status": new_status}

    except Exception:
        if unloading_session_completed and unloading_session:
            await db.unloading_sessions.update_one(
                {"id": unloading_session["id"], "operationId": operation_id},
                {"$set": {"status": "BERJALAN"}, "$unset": {"completedAt": "", "completedBy": "", "operationId": ""}},
            )
        await db.transactions.delete_many({"operation_id": operation_id})
        for change in reversed(stack_changes):
            try:
                await decrease_stack_allocation(change["productId"], change["stackCode"], change["qty"], "Sistem (rollback penerimaan)")
            except Exception:
                pass
        for change in reversed(stock_changes):
            rollback = {"$inc": {change["field"]: -change["qty"], f"channelStock.{change['channel']}.{change['field']}": -change["qty"]}}
            if change["expChanged"]:
                rollback["$set"] = {"exp": change["previousExp"]}
            await db.products.update_one({"id": change["productId"]}, rollback)
        raise

    return {
        "transactions": txns,
        "operationId": operation_id,
        "purchaseOrder": updated_po,
        "message": "Penerimaan stok berhasil disimpan",
    }


@router.post("/stock-damage-discoveries")
async def record_stock_damage_discovery(body: DamageDiscoveryInput, user: dict = Depends(require_write)):
    """Pindahkan stok baik yang ditemukan rusak ke saldo rusak tanpa menghapus jejak asal."""
    product = await db.products.find_one({"id": body.productId}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    _validate_pack_qty(product, float(body.qty))
    stack_code = body.stackCode.strip().upper()
    allocation = await db.stack_allocations.find_one({"productId": body.productId, "stackCode": stack_code}, {"_id": 0, "primaryQty": 1})
    if not allocation or float(allocation.get("primaryQty", 0) or 0) + 1e-9 < float(body.qty):
        raise HTTPException(status_code=400, detail=f"Stok baik pada tumpukan {stack_code} tidak mencukupi")
    channel = normalize_channel(body.channel, normalize_channel(product.get("channel")))
    await ensure_channel_stock(product)
    if float(product.get("stock", 0) or 0) + 1e-9 < float(body.qty) or float(product.get("channelStock", {}).get(channel, {}).get("stock", 0) or 0) + 1e-9 < float(body.qty):
        raise HTTPException(status_code=400, detail="Saldo stok baik tidak mencukupi untuk dipindahkan menjadi stok rusak")

    operation_id, time = new_id(), now_iso()
    await db.products.update_one({"id": product["id"]}, {"$inc": {"stock": -float(body.qty), "damaged": float(body.qty), f"channelStock.{channel}.stock": -float(body.qty), f"channelStock.{channel}.damaged": float(body.qty)}})
    try:
        await decrease_stack_allocation(product["id"], stack_code, float(body.qty), user.get("name", ""))
        transaction = {"id": new_id(), "operation_id": operation_id, "time": time, "ref": body.referenceNo.strip() or f"TR-{time[:10].replace('-', '')}", "type": "PENYESUAIAN", "document_type": "TEMUAN_RUSAK", "kondisi": "RUSAK", "product": product.get("name", ""), "sku": product.get("sku", ""), "change": 0, "good_change": -float(body.qty), "damaged_change": float(body.qty), "unit": product.get("unit", ""), "weight": float(product.get("weight", 0) or 0), "total_weight": float(product.get("weight", 0) or 0) * float(body.qty), "secondary": product.get("secondary", ""), "secondaryQty": float(product.get("secondaryQty", 0) or 0), "channel": channel, "stackCode": stack_code, "operator": user.get("name", ""), "cause": body.cause.strip(), "keterangan": body.note.strip()}
        await db.transactions.insert_one(dict(transaction))
    except Exception:
        try:
            await allocate_stock_to_stack(product, stack_code, float(body.qty), "Sistem (rollback temuan rusak)")
        except Exception:
            pass
        await db.transactions.delete_many({"operation_id": operation_id})
        await db.products.update_one({"id": product["id"]}, {"$inc": {"stock": float(body.qty), "damaged": -float(body.qty), f"channelStock.{channel}.stock": float(body.qty), f"channelStock.{channel}.damaged": -float(body.qty)}})
        raise
    return {"operationId": operation_id, "transaction": transaction}


@router.get("/supplier-returns")
async def list_supplier_returns(user: dict = Depends(get_current_user)):
    return await db.supplier_returns.find({}, {"_id": 0}).sort("created_at", -1).to_list(2000)


@router.post("/supplier-returns")
async def create_supplier_return(body: SupplierReturnInput, user: dict = Depends(require_write)):
    product = await db.products.find_one({"id": body.productId}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    _validate_pack_qty(product, float(body.qty))
    channel = normalize_channel(body.channel, normalize_channel(product.get("channel")))
    await ensure_channel_stock(product)
    if float(product.get("damaged", 0) or 0) + 1e-9 < float(body.qty) or float(product.get("channelStock", {}).get(channel, {}).get("damaged", 0) or 0) + 1e-9 < float(body.qty):
        raise HTTPException(status_code=400, detail="Stok rusak tidak mencukupi untuk diretur ke pemasok")
    source = None
    if body.sourceDamageOperationId:
        source = await db.transactions.find_one({"operation_id": body.sourceDamageOperationId, "product": product.get("name", ""), "kondisi": "RUSAK"}, {"_id": 0})
        if not source:
            raise HTTPException(status_code=400, detail="Sumber barang rusak tidak ditemukan")
        source_qty = float(source.get("damaged_change", source.get("change", 0)) or 0)
        used = await db.supplier_returns.aggregate([{"$match": {"source_damage_operation_id": body.sourceDamageOperationId, "product_id": body.productId}}, {"$group": {"_id": None, "total": {"$sum": "$qty"}}}]).to_list(1)
        if float(body.qty) > source_qty - float(used[0]["total"] if used else 0) + 1e-9:
            raise HTTPException(status_code=400, detail="Jumlah retur melebihi barang rusak pada sumber yang dipilih")
    now = now_iso()
    return_no = body.returnNo.strip() or f"RP-{operational_now().strftime('%Y%m%d')}-{await next_sequence('supplier-return', 0):04d}"
    if await db.supplier_returns.find_one({"return_no": return_no}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="Nomor retur pemasok sudah digunakan")
    doc = {"id": new_id(), "return_no": return_no, "created_at": now, "product_id": body.productId, "product": product.get("name", ""), "sku": product.get("sku", ""), "qty": float(body.qty), "unit": product.get("unit", ""), "channel": channel, "supplier": body.supplier.strip(), "po_no": body.poNo.strip(), "source_damage_operation_id": body.sourceDamageOperationId, "source_stack_code": body.sourceStackCode.strip().upper() or (source or {}).get("stackCode", ""), "note": body.note.strip(), "status": "MENUNGGU_PENGGANTIAN", "replacement_qty": 0.0, "created_by": user.get("name", "")}
    await db.products.update_one({"id": body.productId}, {"$inc": {"damaged": -float(body.qty), f"channelStock.{channel}.damaged": -float(body.qty)}})
    try:
        await db.supplier_returns.insert_one(dict(doc))
        await db.transactions.insert_one({"id": new_id(), "operation_id": doc["id"], "time": now, "ref": return_no, "type": "KELUAR", "document_type": "RETUR_PEMASOK", "parent_document": body.poNo.strip() or body.sourceDamageOperationId, "kondisi": "RUSAK", "product": doc["product"], "sku": doc["sku"], "change": -doc["qty"], "damaged_change": -doc["qty"], "unit": doc["unit"], "channel": channel, "stackCode": doc["source_stack_code"], "penerima": doc["supplier"], "operator": user.get("name", ""), "keterangan": doc["note"]})
    except Exception:
        await db.transactions.delete_many({"operation_id": doc["id"]})
        await db.supplier_returns.delete_many({"id": doc["id"]})
        await db.products.update_one({"id": body.productId}, {"$inc": {"damaged": float(body.qty), f"channelStock.{channel}.damaged": float(body.qty)}})
        raise
    return doc


@router.post("/supplier-returns/{return_id}/replacement")
async def receive_supplier_replacement(return_id: str, body: SupplierReplacementInput, user: dict = Depends(require_write)):
    claim = await db.supplier_returns.find_one({"id": return_id}, {"_id": 0})
    if not claim:
        raise HTTPException(status_code=404, detail="Retur pemasok tidak ditemukan")
    remaining = float(claim.get("qty", 0) or 0) - float(claim.get("replacement_qty", 0) or 0)
    if float(body.qty) > remaining + 1e-9:
        raise HTTPException(status_code=400, detail=f"Penggantian melebihi sisa retur ({remaining:g} {claim.get('unit', '')})")
    product = await db.products.find_one({"id": claim["product_id"]}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Produk pengganti tidak ditemukan")
    _validate_pack_qty(product, float(body.qty))
    stack_code = body.stackCode.strip().upper()
    from backend.stack_allocations import valid_stack_codes
    if stack_code not in await valid_stack_codes():
        raise HTTPException(status_code=400, detail="Tumpukan pengganti tidak valid")
    now = now_iso(); channel = claim.get("channel", normalize_channel(product.get("channel")))
    await ensure_channel_stock(product)
    operation_id = new_id()
    stack_added = False
    claim_updated = False
    replacement_qty = float(claim.get("replacement_qty", 0) or 0) + float(body.qty)
    status = "SELESAI_DIGANTI" if replacement_qty + 1e-9 >= float(claim["qty"]) else "DIGANTI_SEBAGIAN"
    await db.products.update_one({"id": product["id"]}, {"$inc": {"stock": float(body.qty), f"channelStock.{channel}.stock": float(body.qty)}})
    try:
        await allocate_stock_to_stack(product, stack_code, float(body.qty), user.get("name", ""))
        stack_added = True
        result = await db.supplier_returns.update_one(
            {"id": return_id, "replacement_qty": float(claim.get("replacement_qty", 0) or 0)},
            {"$set": {"replacement_qty": replacement_qty, "status": status, "replacement_at": now, "replacement_note": body.note.strip(), "replacement_reference": body.referenceNo.strip()}},
        )
        if result.matched_count == 0:
            raise HTTPException(status_code=409, detail="Data retur pemasok berubah. Muat ulang lalu coba kembali.")
        claim_updated = True
        await db.transactions.insert_one({"id": new_id(), "operation_id": operation_id, "time": now, "ref": body.referenceNo.strip() or f"PG-{claim['return_no']}", "type": "MASUK", "document_type": "PENGGANTIAN_PEMASOK", "parent_document": claim["return_no"], "kondisi": "BAIK", "product": product.get("name", ""), "sku": product.get("sku", ""), "change": float(body.qty), "good_change": float(body.qty), "unit": product.get("unit", ""), "channel": channel, "stackCode": stack_code, "penerima": claim.get("supplier", ""), "operator": user.get("name", ""), "keterangan": body.note.strip()})
    except Exception:
        await db.transactions.delete_many({"operation_id": operation_id})
        if claim_updated:
            await db.supplier_returns.replace_one({"id": return_id, "replacement_qty": replacement_qty}, dict(claim), upsert=False)
        if stack_added:
            try:
                await decrease_stack_allocation(product["id"], stack_code, float(body.qty), "Sistem (rollback penggantian pemasok)")
            except Exception:
                pass
        await db.products.update_one({"id": product["id"]}, {"$inc": {"stock": -float(body.qty), f"channelStock.{channel}.stock": -float(body.qty)}})
        raise
    return await db.supplier_returns.find_one({"id": return_id}, {"_id": 0})


@router.post("/import/master-csv")
async def import_master_csv(file: UploadFile = File(...), user: dict = Depends(require_master_write)):
    raw = await file.read()
    if (file.filename or "").lower().endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                raise HTTPException(status_code=400, detail="File XLSX kosong")
            headers = [str(value or "").strip().lower() for value in values[0]]
            reader = [dict(zip(headers, ["" if value is None else str(value) for value in row])) for row in values[1:]]
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=400, detail="File XLSX tidak dapat dibaca") from exc
    else:
        raise HTTPException(status_code=400, detail="Gunakan file XLSX (.xlsx) dari template aplikasi")
    inserted = 0
    updated = 0
    supplier_names = set()

    for row in reader:
        sku = (row.get("sku") or "").strip().strip("[]")
        name = (row.get("nama") or "").strip()
        if not sku or not name:
            continue

        master = {
            "name": name,
            "sku": sku,
            "category": (row.get("kategori") or "").strip(),
            "cost": _number(row.get("harga_beli"), 0),
            "location": (row.get("lokasi") or "").strip(),
            "supplier": (row.get("supplier") or "").strip(),
            "min": _number(row.get("stok_minimum"), 0),
            "unit": (row.get("satuan") or "Pcs").strip() or "Pcs",
            "weight": _number(row.get("berat_unit"), 0),
            "measureUnit": (row.get("satuan_kuantum") or "kg").strip().lower() or "kg",
            "secondary": (row.get("kemasan_sekunder") or "").strip(),
            "secondaryQty": _number(row.get("isi_kemasan_sekunder"), 0),
            "channel": normalize_channel(row.get("saluran") or row.get("channel")),
        }
        if master["measureUnit"] not in {"kg", "liter", "pcs"}:
            raise HTTPException(status_code=400, detail=f"Satuan kuantum SKU {sku} harus kg, liter, atau pcs")
        fee_columns = {"biaya_muat_buruh":"loadingFeeLabor","biaya_muat_uh":"loadingFeeDaily","biaya_muat_gudang":"loadingFeeWarehouse","lembur_muat_buruh":"loadingOvertimeLabor","lembur_muat_uh":"loadingOvertimeDaily","lembur_muat_gudang":"loadingOvertimeWarehouse","libur_muat_buruh":"loadingHolidayLabor","libur_muat_uh":"loadingHolidayDaily","libur_muat_gudang":"loadingHolidayWarehouse","libur_sore_muat_buruh":"loadingHolidayOvertimeLabor","libur_sore_muat_uh":"loadingHolidayOvertimeDaily","libur_sore_muat_gudang":"loadingHolidayOvertimeWarehouse","biaya_bongkar_buruh":"unloadingFeeLabor","biaya_bongkar_uh":"unloadingFeeDaily","biaya_bongkar_gudang":"unloadingFeeWarehouse","lembur_bongkar_buruh":"unloadingOvertimeLabor","lembur_bongkar_uh":"unloadingOvertimeDaily","lembur_bongkar_gudang":"unloadingOvertimeWarehouse","libur_bongkar_buruh":"unloadingHolidayLabor","libur_bongkar_uh":"unloadingHolidayDaily","libur_bongkar_gudang":"unloadingHolidayWarehouse","libur_sore_bongkar_buruh":"unloadingHolidayOvertimeLabor","libur_sore_bongkar_uh":"unloadingHolidayOvertimeDaily","libur_sore_bongkar_gudang":"unloadingHolidayOvertimeWarehouse"}
        for column, key in fee_columns.items(): master[key] = _number(row.get(column), 0)
        master["loadingFeeChargeMode"] = (row.get("tagihan_muat") or "TIDAK_ADA").strip().upper()
        master["unloadingFeeChargeMode"] = (row.get("tagihan_bongkar") or "TIDAK_ADA").strip().upper()
        if master["secondary"] and master["secondaryQty"] <= 0:
            raise HTTPException(
                status_code=400,
                detail=f"Isi kemasan sekunder untuk SKU {sku} harus lebih dari 0",
            )
        if master["secondaryQty"] > 0 and not master["secondary"]:
            raise HTTPException(status_code=400, detail=f"Kemasan sekunder untuk SKU {sku} wajib diisi")
        if master["secondaryQty"] > 0 and abs(master["secondaryQty"] - round(master["secondaryQty"])) > 1e-6:
            raise HTTPException(status_code=400, detail=f"Isi kemasan sekunder untuk SKU {sku} harus berupa pack utuh")
        if master["supplier"]:
            supplier_names.add(master["supplier"])

        existing = await db.products.find_one({"sku": sku}, {"_id": 0, "id": 1})
        if existing:
            # Import master tidak pernah mengubah stock, damaged, atau expired.
            await db.products.update_one({"sku": sku}, {"$set": master})
            updated += 1
        else:
            doc = {
                **master,
                "id": new_id(),
                "stock": 0,
                "damaged": 0,
                "exp": "",
            }
            await db.products.insert_one(doc)
            inserted += 1

    if inserted == 0 and updated == 0:
        raise HTTPException(status_code=400, detail="File tidak berisi master SKU yang valid")

    existing_suppliers = {
        doc["name"]
        for doc in await db.suppliers.find({}, {"_id": 0, "name": 1}).to_list(1000)
    }
    for name in sorted(supplier_names - existing_suppliers):
        await db.suppliers.insert_one({
            "id": new_id(), "name": name, "pic": "", "phone": "",
            "email": "", "address": "", "category": "",
        })

    return {"inserted": inserted, "updated": updated}


@router.get("/export/master-template.xlsx")
async def export_master_template(user: dict = Depends(get_current_user)):
    headers = ["sku","nama","kategori","saluran","satuan","satuan_kuantum","berat_unit","kemasan_sekunder","isi_kemasan_sekunder","harga_beli","supplier","lokasi","stok_minimum","tagihan_muat","biaya_muat_buruh","biaya_muat_uh","biaya_muat_gudang","lembur_muat_buruh","lembur_muat_uh","lembur_muat_gudang","libur_muat_buruh","libur_muat_uh","libur_muat_gudang","libur_sore_muat_buruh","libur_sore_muat_uh","libur_sore_muat_gudang","tagihan_bongkar","biaya_bongkar_buruh","biaya_bongkar_uh","biaya_bongkar_gudang","lembur_bongkar_buruh","lembur_bongkar_uh","lembur_bongkar_gudang","libur_bongkar_buruh","libur_bongkar_uh","libur_bongkar_gudang","libur_sore_bongkar_buruh","libur_sore_bongkar_uh","libur_sore_bongkar_gudang"]
    rows = [
        ["B0010001X","CONTOH BERAS MEDIUM 5 KG","Beras","PSO","Pack","kg",5,"Karung",8,0,"Nama Supplier","GBB 17",0,"PENGAMBIL",410,10,60,50,5,0,75,10,0,25,5,0,"TIDAK_ADA",0,0,0,0,0,0,0,0,0,0,0,0],
        ["B0100152X","CONTOH MINYAK 2 L","Minyak","KOM","Botol","liter",2,"Dus",6,0,"Nama Supplier","GBB 18",0,"PENGAMBIL",480,10,60,50,5,0,75,10,0,25,5,0,"TIDAK_ADA",0,0,0,0,0,0,0,0,0,0,0,0],
    ]
    output = build_xlsx(headers, rows, "Master Produk")
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="template_import_master_produk.xlsx"'})


@router.get("/export/transactions-v2.xlsx")
async def export_transactions_with_expiry(user: dict = Depends(get_current_user)):
    now = operational_now()
    start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    if start.month == 12:
        end = start.replace(year=start.year + 1, month=1)
    else:
        end = start.replace(month=start.month + 1)

    transactions = await db.transactions.find(
        {
            "time": {
                "$gte": start.astimezone(timezone.utc).isoformat(),
                "$lt": end.astimezone(timezone.utc).isoformat(),
            }
        },
        {"_id": 0},
    ).sort("time", 1).to_list(10000)

    headers = [
        "Waktu", "No. Referensi", "No. PO", "Antrian", "Tipe", "Kondisi",
        "Produk", "SKU", "Jumlah", "Satuan", "Tanggal Kedaluwarsa",
        "Berat Total (kg)", "Kemasan Sekunder", "Isi/Kemasan Sekunder",
        "Pihak Terkait", "No. Polisi", "Jenis Dokumen", "Dokumen Induk",
        "Jumlah Diselesaikan SO", "Good Kembali", "Damage Kembali", "Dicatat Oleh", "Keterangan",
    ]
    rows = [
        [
            t.get("time", ""), t.get("ref", ""), t.get("po_no", ""), t.get("antrian", ""),
            t.get("type", ""), t.get("kondisi", ""), t.get("product", ""), t.get("sku", ""),
            t.get("change", 0), t.get("unit", ""), t.get("exp", ""),
            t.get("total_weight", abs(t.get("change", 0) or 0) * (t.get("weight", 0) or 0)),
            t.get("secondary", ""), t.get("secondaryQty", 0), t.get("penerima", ""),
            t.get("polisi", ""), t.get("document_type", ""), t.get("parent_document", ""),
            t.get("settled_qty", ""), t.get("good_change", ""), t.get("damaged_change", ""),
            t.get("operator", ""), t.get("keterangan", ""),
        ]
        for t in transactions
    ]
    output = build_xlsx(headers, rows, "Riwayat Transaksi")
    filename = f"riwayat_transaksi_{start.strftime('%Y_%m')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
