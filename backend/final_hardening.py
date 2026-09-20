from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from openpyxl.styles import Alignment, Font, PatternFill

from backend.server import (
    JAKARTA_TZ,
    db,
    get_current_user,
    new_id,
    normalize_channel,
    operational_now,
    require_master_write,
)
from backend.role_four_config import has_role_permission, role_destination
from backend.so_monitoring import build_so_monitoring

router = APIRouter(prefix="/api")
EPS = 1e-9
MAX_IMPORT_BYTES = 5 * 1024 * 1024

MASTER_HEADERS = [
    "SKU", "Nama", "Kategori", "Saluran", "Satuan Primer", "Kuantum/Unit",
    "Satuan Kuantum", "Kemasan Sekunder", "Isi Kemasan Sekunder", "Supplier",
    "Lokasi Default", "Stok Minimum",
    "Biaya Muat Buruh", "Biaya Muat Harian", "Biaya Muat Gudang", "Mode Biaya Muat",
    "Biaya Bongkar Buruh", "Biaya Bongkar Harian", "Biaya Bongkar Gudang", "Mode Biaya Bongkar",
]


def _n(value, default=0.0) -> float:
    if value is None or value == "":
        return float(default)
    try:
        if isinstance(value, str):
            value = value.strip().replace(" ", "").replace(",", ".")
        return float(value)
    except (TypeError, ValueError):
        return float(default)


def _text(value) -> str:
    return str(value or "").strip()


def _date_range(start: str, end: str) -> tuple[str, str, str, str]:
    today = operational_now().date()
    try:
        start_date = datetime.strptime(start, "%Y-%m-%d").date() if start else today.replace(day=1)
        end_date = datetime.strptime(end, "%Y-%m-%d").date() if end else today
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Tanggal laporan harus YYYY-MM-DD") from exc
    if end_date < start_date:
        raise HTTPException(status_code=400, detail="Tanggal akhir tidak boleh lebih kecil dari tanggal awal")
    if (end_date - start_date).days > 366:
        raise HTTPException(status_code=400, detail="Rentang laporan maksimal 366 hari")
    start_dt = datetime.combine(start_date, datetime.min.time(), JAKARTA_TZ).astimezone(timezone.utc)
    end_dt = datetime.combine(end_date + timedelta(days=1), datetime.min.time(), JAKARTA_TZ).astimezone(timezone.utc)
    return start_date.isoformat(), end_date.isoformat(), start_dt.isoformat(), end_dt.isoformat()


def _style_sheet(ws):
    if ws.max_row:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        letter = column[0].column_letter
        max_len = max((len(str(cell.value or "")) for cell in column), default=0)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)


def _append_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]]):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_sheet(ws)
    return ws


@router.get("/dashboard-operations")
async def dashboard_operations(user: dict = Depends(get_current_user)):
    scoped = role_destination(user.get("role"))
    if scoped:
        active_opnames = await db.consignment_opnames.count_documents({
            "destination": scoped,
            "status": {"$in": ["DRAFT", "SUBMITTED"]},
        })
        damaged_opnames = await db.consignment_damaged_opnames.count_documents({
            "destination": scoped,
            "status": {"$in": ["DRAFT", "SUBMITTED"]},
        })
        return {
            "scope": scoped,
            "consignmentOpnamesPending": active_opnames + damaged_opnames,
        }

    so = await build_so_monitoring()
    loads = await db.outbound_loads.find(
        {"status": {"$in": ["Menunggu", "Sedang Dimuat"]}},
        {"_id": 0, "id": 1, "status": 1, "bon_no": 1, "antrian": 1, "documents": 1, "items": 1},
    ).to_list(5000)
    loading_unpaid = None
    if has_role_permission(user.get("role"), "costView"):
        loading_unpaid = await db.outbound_loads.count_documents({
            "status": "Selesai",
            "loading_cost.chargeable": {"$gt": 0},
            "loading_fee_payment_status": {"$in": ["BELUM_DIBAYAR", "SEBAGIAN", "", None]},
        })
    main_opnames = await db.stock_opnames.count_documents({"status": {"$in": ["DRAFT", "SUBMITTED"]}})
    cons_opnames = await db.consignment_opnames.count_documents({"status": {"$in": ["DRAFT", "SUBMITTED"]}})
    damaged_opnames = await db.consignment_damaged_opnames.count_documents({"status": {"$in": ["DRAFT", "SUBMITTED"]}})
    arrangement_pending = await db.stack_allocations.count_documents({
        "primaryQty": {"$gt": 0},
        "arrangementAdjusted": True,
    })
    integrity_open = await db.operational_postcommit_issues.count_documents({"status": "OPEN"})

    pending_doc_loads = await db.outbound_loads.find(
        {"status": "Selesai", "document_type": {"$in": ["CT", "MEMO", "ND"]}},
        {"_id": 0, "id": 1, "document_type": 1, "ref": 1, "document_status": 1, "document_links": 1, "items": 1},
    ).to_list(10000)
    pending_documents = 0
    for load in pending_doc_loads:
        if not load.get("document_links"):
            pending_documents += 1
        elif str(load.get("document_status") or "").upper() not in {"SELESAI", "SETTLED"}:
            pending_documents += 1

    return {
        "scope": "MAIN",
        "soOutstanding": int((so.get("summary") or {}).get("outstanding", 0) or 0),
        "soPartial": int((so.get("summary") or {}).get("partial", 0) or 0),
        "activeQueue": len(loads),
        "activeQueueItems": sum(len(load.get("items") or []) for load in loads),
        "pendingDocuments": pending_documents,
        "loadingPaymentsPending": loading_unpaid,
        "opnamesPending": main_opnames + cons_opnames + damaged_opnames,
        "arrangementPending": arrangement_pending,
        "postCommitOpen": integrity_open if has_role_permission(user.get("role"), "masterWrite") else None,
    }


@router.get("/product-catalog")
async def product_catalog(user: dict = Depends(get_current_user)):
    # Katalog aman untuk role Bazar/E-commerce: tidak mengekspos saldo utama, biaya,
    # supplier, ataupun nilai persediaan.
    rows = await db.products.find({}, {
        "_id": 0, "id": 1, "sku": 1, "name": 1, "category": 1, "channel": 1,
        "unit": 1, "weight": 1, "measureUnit": 1, "secondary": 1, "secondaryQty": 1,
    }).sort([("name", 1), ("sku", 1)]).to_list(20000)
    return rows


@router.get("/reports/operational.xlsx")
async def operational_report(
    start: str = Query(default=""),
    end: str = Query(default=""),
    user: dict = Depends(get_current_user),
):
    if role_destination(user.get("role")):
        raise HTTPException(status_code=403, detail="Laporan Gudang Utama tidak tersedia untuk role Bazar/E-commerce")
    start_label, end_label, start_utc, end_utc = _date_range(start, end)
    date_query = {"$gte": start_utc, "$lt": end_utc}

    txns = await db.transactions.find({"time": date_query}, {"_id": 0}).sort("time", 1).to_list(100000)
    loads = await db.outbound_loads.find(
        {"created_at": date_query},
        {"_id": 0},
    ).sort("created_at", 1).to_list(30000)
    opnames = await db.stock_opnames.find(
        {"createdAt": date_query},
        {"_id": 0},
    ).sort("createdAt", 1).to_list(20000)
    cons_moves = await db.consignment_movements.find(
        {"time": date_query},
        {"_id": 0},
    ).sort("time", 1).to_list(100000)
    stack_history = await db.stack_history.find(
        {"time": date_query},
        {"_id": 0},
    ).sort("time", 1).to_list(100000)
    stack_snapshot = await db.stack_allocations.find(
        {},
        {"_id": 0},
    ).sort([("stackCode", 1), ("productName", 1)]).to_list(50000)

    wb = Workbook()
    wb.remove(wb.active)

    _append_sheet(wb, "Transaksi", [
        "Waktu", "Tipe", "Kondisi", "Referensi", "Dokumen", "SKU", "Produk",
        "Perubahan", "Satuan", "Tumpukan", "Pihak", "Nopol", "Petugas",
    ], [[
        t.get("time", ""), t.get("type", ""), t.get("kondisi", ""), t.get("ref", ""),
        t.get("documentNo", ""), t.get("sku", ""), t.get("product", ""), t.get("change", 0),
        t.get("unit", ""), t.get("stackCode", ""), t.get("penerima", ""), t.get("polisi", ""),
        t.get("operator", ""),
    ] for t in txns])

    load_rows = []
    for load in loads:
        for item in load.get("items") or []:
            load_rows.append([
                load.get("operational_date", ""), load.get("status", ""), load.get("bon_no", ""),
                load.get("antrian", ""), item.get("documentNo") or load.get("ref", ""),
                item.get("sku", ""), item.get("name", ""), item.get("qty", 0), item.get("unit", ""),
                item.get("stackCode", ""), load.get("party", ""), load.get("polisi", ""),
                load.get("surat_jalan_no", ""), _n((item.get("loadingFee") or {}).get("total")),
            ])
    _append_sheet(wb, "Pengeluaran", [
        "Tanggal", "Status", "Bon Muat", "Antrian", "Dokumen", "SKU", "Produk",
        "Qty", "Satuan", "Tumpukan", "Penerima", "Nopol", "Surat Jalan", "Biaya Muat",
    ], load_rows)

    _append_sheet(wb, "Stock Opname", [
        "Waktu", "Nomor", "Gudang", "Status", "Petugas", "Disetujui Oleh",
    ], [[
        row.get("createdAt", ""), row.get("no", ""), row.get("warehouse", ""),
        row.get("status", ""), row.get("operator", ""), row.get("approvedBy", ""),
    ] for row in opnames])

    _append_sheet(wb, "Mutasi Tumpukan", [
        "Waktu", "Aksi", "Tumpukan", "SKU", "Produk", "Saldo Setelah", "Satuan", "Petugas",
    ], [[
        row.get("time", ""), row.get("action", ""), row.get("stackCode", ""),
        (row.get("allocation") or row.get("after") or {}).get("sku", ""),
        (row.get("allocation") or row.get("after") or {}).get("productName", ""),
        (row.get("allocation") or row.get("after") or {}).get("primaryQty", 0),
        (row.get("allocation") or row.get("after") or {}).get("unit", ""),
        row.get("operator", ""),
    ] for row in stack_history])

    _append_sheet(wb, "Snapshot Tumpukan", [
        "Tumpukan", "Gudang", "Zona", "SKU", "Produk", "Saldo Primer", "Satuan",
        "Kuantum/Unit", "Satuan Kuantum", "Kemasan Sekunder", "Isi Sekunder", "Susunan Perlu Update",
    ], [[
        row.get("stackCode", ""), row.get("warehouse", ""), row.get("zone", ""), row.get("sku", ""),
        row.get("productName", ""), row.get("primaryQty", 0), row.get("unit", ""), row.get("weight", 0),
        row.get("measureUnit", ""), row.get("secondary", ""), row.get("secondaryQty", 0),
        "YA" if row.get("arrangementAdjusted") else "TIDAK",
    ] for row in stack_snapshot])

    _append_sheet(wb, "Bazar-Ecom", [
        "Waktu", "Lokasi", "Tipe", "Referensi", "SKU", "Produk", "Delta", "Satuan", "Petugas",
    ], [[
        row.get("time", ""), row.get("destination", ""), row.get("movementType", ""),
        row.get("referenceNo", ""), row.get("sku", ""), row.get("name", ""), row.get("delta", 0),
        row.get("unit", ""), row.get("operator", ""),
    ] for row in cons_moves])

    so_monitoring = await build_so_monitoring()
    so_rows = []
    for record in so_monitoring.get("records", []):
        for item in record.get("items", []):
            so_rows.append([
                record.get("documentNo", ""), record.get("status", ""), record.get("party", ""),
                item.get("sku", ""), item.get("name", ""), item.get("orderedQty"),
                item.get("completedQty", 0), item.get("reservedQty", 0), item.get("remainingQty"),
                item.get("unit", ""), record.get("lastActivity", ""),
            ])
    _append_sheet(wb, "Outstanding SO", [
        "Nomor SO", "Status", "Penerima", "SKU", "Produk", "Kuantum SO", "Selesai",
        "Reservasi", "Sisa Dijadwalkan", "Satuan", "Aktivitas Terakhir",
    ], so_rows)

    if has_role_permission(user.get("role"), "costView"):
        cost_rows = []
        for load in loads:
            cost = load.get("loading_cost") or {}
            if _n(cost.get("total")) <= EPS:
                continue
            cost_rows.append([
                load.get("operational_date", ""), load.get("bon_no", ""), load.get("ref", ""),
                cost.get("labor", 0), cost.get("daily", 0), cost.get("warehouse", 0), cost.get("total", 0),
                cost.get("chargeable", 0), load.get("loading_fee_payment_total", 0),
                load.get("loading_fee_payment_status", ""),
            ])
        _append_sheet(wb, "Biaya Muat", [
            "Tanggal", "Bon Muat", "Referensi", "Buruh", "Harian", "Gudang", "Total",
            "Ditagihkan", "Dibayar", "Status Bayar",
        ], cost_rows)

    info = wb.create_sheet("Info", 0)
    info.append(["LAPORAN OPERASIONAL PEPEG"])
    info.append(["Periode", f"{start_label} s.d. {end_label}"])
    info.append(["Dicetak", operational_now().isoformat()])
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 42

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"laporan_operasional_{start_label}_{end_label}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/reports/operational.pdf")
async def operational_report_pdf(
    start: str = Query(default=""),
    end: str = Query(default=""),
    user: dict = Depends(get_current_user),
):
    if role_destination(user.get("role")):
        raise HTTPException(status_code=403, detail="Laporan Gudang Utama tidak tersedia untuk role Bazar/E-commerce")
    start_label, end_label, start_utc, end_utc = _date_range(start, end)
    date_query = {"$gte": start_utc, "$lt": end_utc}

    tx_count = await db.transactions.count_documents({"time": date_query})
    inbound_count = await db.transactions.count_documents({"time": date_query, "type": "MASUK"})
    outbound_count = await db.transactions.count_documents({"time": date_query, "type": "KELUAR"})
    loads = await db.outbound_loads.find(
        {"created_at": date_query},
        {"_id": 0, "status": 1, "bon_no": 1, "ref": 1, "operational_date": 1, "surat_jalan_no": 1, "items": 1, "loading_cost": 1},
    ).sort("created_at", 1).to_list(30000)
    opnames_pending = await db.stock_opnames.count_documents({"status": {"$in": ["DRAFT", "SUBMITTED"]}})
    arrangement_pending = await db.stack_allocations.count_documents({"primaryQty": {"$gt": 0}, "arrangementAdjusted": True})
    so = await build_so_monitoring()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=12*mm, rightMargin=12*mm, topMargin=12*mm, bottomMargin=12*mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("title", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=16, leading=20, alignment=TA_CENTER)
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=7.5, leading=10)
    story = [
        Paragraph("LAPORAN OPERASIONAL PEPEG", title),
        Paragraph(f"Periode {start_label} s.d. {end_label}", ParagraphStyle("sub", parent=small, alignment=TA_CENTER, fontSize=9, leading=12)),
        Spacer(1, 6*mm),
    ]
    summary_data = [
        ["Transaksi", tx_count, "Penerimaan", inbound_count, "Pengeluaran", outbound_count],
        ["SO Outstanding", int((so.get("summary") or {}).get("outstanding", 0) or 0), "Opname Pending", opnames_pending, "Susunan Perlu Update", arrangement_pending],
    ]
    summary_table = Table(summary_data, colWidths=[35*mm, 22*mm, 35*mm, 22*mm, 40*mm, 22*mm])
    summary_table.setStyle(TableStyle([
        ("GRID",(0,0),(-1,-1),.4,colors.HexColor("#94a3b8")),
        ("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#f8fafc")),
        ("FONTNAME",(0,0),(-1,-1),"Helvetica-Bold"),
        ("ALIGN",(1,0),(1,-1),"RIGHT"),("ALIGN",(3,0),(3,-1),"RIGHT"),("ALIGN",(5,0),(5,-1),"RIGHT"),
        ("FONTSIZE",(0,0),(-1,-1),8),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5),
    ]))
    story += [summary_table, Spacer(1, 6*mm), Paragraph("RINGKASAN PEMUATAN", ParagraphStyle("h2", parent=title, fontSize=11, leading=14, alignment=0))]
    data = [["Tanggal","Bon Muat","Referensi","Status","Produk / Qty","Surat Jalan"]]
    for load in loads[:200]:
        item_text = "; ".join(f"{item.get('name','')}: {_n(item.get('qty')):g} {item.get('unit','')}" for item in load.get("items") or [])
        data.append([
            load.get("operational_date",""), load.get("bon_no",""), load.get("ref",""), load.get("status",""),
            Paragraph(item_text[:220] or "-", small), load.get("surat_jalan_no",""),
        ])
    table = Table(data, colWidths=[24*mm,34*mm,40*mm,24*mm,105*mm,44*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1f4e78")),("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),6.5),
        ("GRID",(0,0),(-1,-1),.3,colors.HexColor("#94a3b8")),("VALIGN",(0,0),(-1,-1),"TOP"),
        ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3),
    ]))
    story.append(table)
    if len(loads) > 200:
        story += [Spacer(1, 3*mm), Paragraph(f"Catatan: PDF menampilkan 200 pemuatan pertama dari {len(loads)}. Gunakan XLSX untuk detail lengkap.", small)]
    if has_role_permission(user.get("role"), "costView"):
        total_cost = sum(_n((load.get("loading_cost") or {}).get("total")) for load in loads)
        story += [Spacer(1, 5*mm), Paragraph(f"Total biaya muat pada periode: Rp {total_cost:,.0f}".replace(",", "."), small)]
    story += [Spacer(1, 5*mm), Paragraph(f"Dicetak: {operational_now().strftime('%d-%m-%Y %H:%M WIB')}", small)]
    doc.build(story)
    buffer.seek(0)
    filename = f"laporan_operasional_{start_label}_{end_label}.pdf"
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _master_template() -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Master Produk"
    ws.append(MASTER_HEADERS)
    ws.append([
        "CONTOH-001", "Contoh Produk", "Beras", "KOM", "Pack", 5, "kg",
        "Karung", 8, "Contoh Supplier", "", 0,
        0, 0, 0, "TIDAK_ADA", 0, 0, 0, "TIDAK_ADA",
    ])
    _style_sheet(ws)
    panduan = wb.create_sheet("Panduan")
    rows = [
        ["ATURAN IMPORT MASTER"],
        ["1", "File ini hanya mengubah MASTER. Stok baik/rusak, transaksi, lot, tumpukan, dan kedaluwarsa tidak pernah diimport."],
        ["2", "SKU harus unik. SKU yang sudah ada hanya memperbarui metadata yang aman."],
        ["3", "Saluran hanya PSO atau KOM. Satuan kuantum hanya kg, liter, atau pcs."],
        ["4", "Kemasan sekunder dan isi kemasan tidak dapat diubah bila produk masih dialokasikan pada tumpukan."],
        ["5", "Biaya buruh tidak wajib diisi; kolom kosong dianggap 0."],
    ]
    for row in rows:
        panduan.append(row)
    panduan.column_dimensions["A"].width = 8
    panduan.column_dimensions["B"].width = 110
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@router.get("/export/master-template.xlsx")
async def master_template(user: dict = Depends(require_master_write)):
    return StreamingResponse(
        _master_template(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="template_import_master_produk.xlsx"'},
    )


def _header_map(ws) -> dict[str, int]:
    return {_text(cell.value).lower(): idx for idx, cell in enumerate(ws[1]) if _text(cell.value)}


def _cell(row, headers: dict[str, int], name: str, default=None):
    index = headers.get(name.lower())
    return row[index] if index is not None and index < len(row) else default


def _normalize_import_row(row, headers) -> dict:
    sku = _text(_cell(row, headers, "SKU")).upper()
    name = _text(_cell(row, headers, "Nama"))
    if not sku or not name:
        raise ValueError("SKU dan Nama wajib diisi")
    channel = _text(_cell(row, headers, "Saluran", "KOM")).upper() or "KOM"
    if channel not in {"PSO", "KOM"}:
        raise ValueError("Saluran harus PSO atau KOM")
    measure = _text(_cell(row, headers, "Satuan Kuantum", "kg")).lower() or "kg"
    if measure not in {"kg", "liter", "pcs"}:
        raise ValueError("Satuan Kuantum harus kg, liter, atau pcs")
    secondary = _text(_cell(row, headers, "Kemasan Sekunder"))
    secondary_qty = _n(_cell(row, headers, "Isi Kemasan Sekunder"))
    if (secondary and secondary_qty <= 0) or (secondary_qty > 0 and not secondary):
        raise ValueError("Kemasan Sekunder dan Isi Kemasan Sekunder harus diisi berpasangan")
    if secondary_qty > 0 and abs(secondary_qty - round(secondary_qty)) > 1e-6:
        raise ValueError("Isi Kemasan Sekunder harus bilangan utuh")
    loading_mode = _text(_cell(row, headers, "Mode Biaya Muat", "TIDAK_ADA")).upper() or "TIDAK_ADA"
    unloading_mode = _text(_cell(row, headers, "Mode Biaya Bongkar", "TIDAK_ADA")).upper() or "TIDAK_ADA"
    if loading_mode not in {"PENGAMBIL", "TERMASUK", "TIDAK_ADA"}:
        raise ValueError("Mode Biaya Muat tidak valid")
    if unloading_mode not in {"PENGIRIM", "TERMASUK", "TIDAK_ADA"}:
        raise ValueError("Mode Biaya Bongkar tidak valid")
    return {
        "sku": sku,
        "name": name,
        "category": _text(_cell(row, headers, "Kategori")),
        "channel": channel,
        "unit": _text(_cell(row, headers, "Satuan Primer", "Pcs")) or "Pcs",
        "weight": max(_n(_cell(row, headers, "Kuantum/Unit")), 0),
        "measureUnit": measure,
        "secondary": secondary,
        "secondaryQty": secondary_qty,
        "supplier": _text(_cell(row, headers, "Supplier")),
        "location": _text(_cell(row, headers, "Lokasi Default")),
        "min": max(_n(_cell(row, headers, "Stok Minimum")), 0),
        "loadingFeeLabor": max(_n(_cell(row, headers, "Biaya Muat Buruh")), 0),
        "loadingFeeDaily": max(_n(_cell(row, headers, "Biaya Muat Harian")), 0),
        "loadingFeeWarehouse": max(_n(_cell(row, headers, "Biaya Muat Gudang")), 0),
        "loadingFeeChargeMode": loading_mode,
        "unloadingFeeLabor": max(_n(_cell(row, headers, "Biaya Bongkar Buruh")), 0),
        "unloadingFeeDaily": max(_n(_cell(row, headers, "Biaya Bongkar Harian")), 0),
        "unloadingFeeWarehouse": max(_n(_cell(row, headers, "Biaya Bongkar Gudang")), 0),
        "unloadingFeeChargeMode": unloading_mode,
    }


async def _import_master_xlsx(file: UploadFile) -> dict:
    filename = _text(file.filename).lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Import master hanya menerima file .xlsx")
    content = await file.read(MAX_IMPORT_BYTES + 1)
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="Ukuran file maksimal 5 MB")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="File XLSX tidak dapat dibaca") from exc
    if "Master Produk" not in wb.sheetnames:
        raise HTTPException(status_code=400, detail="Sheet 'Master Produk' tidak ditemukan. Gunakan template PEPEG.")
    ws = wb["Master Produk"]
    headers = _header_map(ws)
    if "sku" not in headers or "nama" not in headers:
        raise HTTPException(status_code=400, detail="Kolom SKU dan Nama wajib tersedia")

    parsed = []
    errors = []
    seen = set()
    for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not any(value not in (None, "") for value in values):
            continue
        try:
            doc = _normalize_import_row(values, headers)
            key = doc["sku"].upper()
            if key in seen:
                raise ValueError("SKU duplikat di dalam file")
            seen.add(key)
            parsed.append((row_no, doc))
        except ValueError as exc:
            errors.append({"row": row_no, "error": str(exc)})
    if errors:
        raise HTTPException(status_code=400, detail={"message": "Import dibatalkan karena ada baris tidak valid", "errors": errors[:50]})
    if not parsed:
        raise HTTPException(status_code=400, detail="Tidak ada data produk yang dapat diimport")

    existing = {
        str(row.get("sku") or "").upper(): row
        for row in await db.products.find({}, {"_id": 0}).to_list(50000)
    }
    settings = await db.settings.find_one({"_id": "app"}, {"_id": 0, "categories": 1}) or {}
    valid_categories = {
        _text(item.get("name")).lower()
        for item in settings.get("categories") or []
        if _text(item.get("name"))
    }
    category_errors = [
        {"row": row_no, "error": f"Kategori {doc.get('category')} belum terdaftar di Pengaturan"}
        for row_no, doc in parsed
        if doc.get("category") and valid_categories and doc.get("category", "").lower() not in valid_categories
    ]
    if category_errors:
        raise HTTPException(status_code=400, detail={"message": "Import dibatalkan karena kategori belum terdaftar", "errors": category_errors[:50]})

    inserted = updated = protected = 0
    suppliers = set()
    for _, original_doc in parsed:
        doc = dict(original_doc)
        key = doc["sku"].upper()
        current = existing.get(key)
        if current:
            has_balance = _n(current.get("stock")) > EPS or _n(current.get("damaged")) > EPS
            if has_balance and doc.get("channel") != str(current.get("channel") or "KOM").upper():
                doc["channel"] = str(current.get("channel") or "KOM").upper()
                protected += 1

            packaging_changed = (
                _n(current.get("secondaryQty")) != _n(doc.get("secondaryQty"))
                or _text(current.get("secondary")) != _text(doc.get("secondary"))
            )
            has_allocation = bool(await db.stack_allocations.find_one({"productId": current["id"]}, {"_id": 1}))
            if packaging_changed and has_allocation:
                doc["secondary"] = current.get("secondary", "")
                doc["secondaryQty"] = _n(current.get("secondaryQty"))
                protected += 1

            # Jangan pernah menyentuh stock, damaged, exp, channelStock, lot, ataupun id.
            await db.products.update_one({"id": current["id"]}, {"$set": doc})
            updated += 1
        else:
            await db.products.insert_one({
                **doc,
                "id": new_id(),
                "stock": 0,
                "damaged": 0,
                "exp": "",
            })
            inserted += 1
        if doc.get("supplier"):
            suppliers.add(doc["supplier"])

    existing_suppliers = {
        str(row.get("name") or "").strip().lower()
        for row in await db.suppliers.find({}, {"_id": 0, "name": 1}).to_list(10000)
    }
    suppliers_added = 0
    for name in sorted(suppliers):
        if name.lower() in existing_suppliers:
            continue
        await db.suppliers.insert_one({
            "id": new_id(), "name": name, "pic": "", "phone": "", "email": "", "address": "", "category": "",
        })
        suppliers_added += 1

    return {
        "inserted": inserted,
        "updated": updated,
        "protectedFields": protected,
        "skipped": 0,
        "suppliersAdded": suppliers_added,
        "rows": len(parsed),
    }


@router.post("/import/master-xlsx")
async def import_master_xlsx(file: UploadFile = File(...), user: dict = Depends(require_master_write)):
    return await _import_master_xlsx(file)


@router.post("/import/master-csv")
async def import_master_compat(file: UploadFile = File(...), user: dict = Depends(require_master_write)):
    # Nama endpoint dipertahankan agar frontend lama tidak putus; format yang sah tetap XLSX.
    return await _import_master_xlsx(file)


@router.get("/role-permission-matrix")
async def role_permission_matrix(user: dict = Depends(get_current_user)):
    if not has_role_permission(user.get("role"), "users"):
        raise HTTPException(status_code=403, detail="Hanya Superadmin yang dapat melihat audit permission")
    roles = ["Administrator", "Kepala Gudang", "Supervisor", "Operator", "QC", "Pemantau"]
    permissions = [
        "mainInventory", "masterWrite", "inbound", "outbound", "costView", "corrections",
        "warehouseApprove", "users", "settings", "bazarView", "bazarOps", "ecomView", "ecomOps",
        "consignmentView", "consignmentHistory",
    ]
    return {
        "roles": roles,
        "permissions": permissions,
        "matrix": {
            role: {permission: bool(has_role_permission(role, permission)) for permission in permissions}
            for role in roles
        },
    }
