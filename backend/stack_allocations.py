import io
import re
from datetime import datetime
from typing import List, Literal, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from backend.server import build_xlsx, db, get_current_user, new_id, now_iso, operational_now, require_write

router = APIRouter(prefix="/api")


async def valid_stack_codes() -> set[str]:
    settings = await db.settings.find_one({"_id": "app"}, {"_id": 0, "warehouses": 1}) or {}
    warehouses = settings.get("warehouses") or [
        *[{"code": str(unit), "zones": [{"code": zone, "count": 4} for zone in ("A", "B", "C")]} for unit in range(17, 25)],
        {"code": "MP1", "zones": [{"code": zone, "count": 8} for zone in ("A", "B")]},
    ]
    return {
        f"{str(warehouse.get('code', '')).upper()}/{str(zone.get('code', '')).upper()}{number:02d}"
        for warehouse in warehouses
        for zone in warehouse.get("zones", [])
        for number in range(1, int(zone.get("count", 0) or 0) + 1)
    }


async def allocate_stock_to_stack(product: dict, stack_code: str, qty: float, operator: str = "") -> None:
    stack_code = stack_code.strip().upper()
    if stack_code not in await valid_stack_codes():
        raise HTTPException(status_code=400, detail="Lokasi tumpukan penerimaan tidak valid")
    per_secondary = float(product.get("secondaryQty", 0) or 0)
    if per_secondary <= 0 or not product.get("secondary"):
        raise HTTPException(status_code=400, detail=f"Kemasan sekunder {product.get('name', '')} belum diatur")
    existing = await db.stack_allocations.find_one({"productId": product["id"], "stackCode": stack_code}, {"_id": 0})
    if existing:
        total = float(existing.get("primaryQty", 0) or 0) + qty
        changes = {
            "primaryQty": total, "secondaryCount": int(total // per_secondary),
            "primaryRemainder": total % per_secondary, "length": 0, "width": 0, "height": 0,
            "arrangementAdjusted": True, "updatedAt": now_iso(),
        }
        await db.stack_allocations.update_one({"id": existing["id"]}, {"$set": changes})
        await record_stack_history("PENERIMAAN_OTOMATIS", {**existing, **changes}, operator or "Sistem")
        return
    allocation = {
        "id": new_id(), "productId": product["id"], "sku": product.get("sku", ""),
        "productName": product.get("name", ""), "unit": product.get("unit", ""),
        "weight": float(product.get("weight", 0) or 0), "measureUnit": product.get("measureUnit", "kg") or "kg", "secondary": product.get("secondary", ""),
        "secondaryQty": per_secondary, "stackCode": stack_code, "warehouse": stack_code.split('/')[0],
        "zone": re.sub(r"\d", "", stack_code.split('/')[1]), "length": 0, "width": 0, "height": 0,
        "secondaryCount": int(qty // per_secondary), "primaryRemainder": qty % per_secondary,
        "primaryQty": qty, "note": f"Penerimaan oleh {operator}" if operator else "Penerimaan",
        "arrangementAdjusted": True, "createdAt": now_iso(), "updatedAt": now_iso(),
    }
    await db.stack_allocations.insert_one(dict(allocation))
    await record_stack_history("PENERIMAAN_OTOMATIS", allocation, operator or "Sistem")


class StackArrangement(BaseModel):
    hamparan: int = Field(ge=1, le=1000)
    kaki: int = Field(ge=1, le=1000)
    height: int = Field(ge=1, le=1000)


class StackAllocationBody(BaseModel):
    productId: str
    stackCode: str
    length: int = Field(ge=1, le=1000)
    width: int = Field(ge=1, le=1000)
    height: int = Field(ge=1, le=1000)
    arrangements: List[StackArrangement] = Field(default_factory=list, max_length=10)
    extraSecondary: int = Field(default=0, ge=0, le=1000000)
    extraPrimary: int = Field(default=0, ge=0, le=1000000)
    note: str = ""


class StackTreatmentBody(BaseModel):
    type: Literal["SPRAYING", "FUMIGASI", "FUMIGASI_SULFUR"]
    warehouse: str
    stackCode: str = ""
    startDate: str
    endDate: str = ""
    note: str = ""


def _parse_treatment_date(value: str, label: str):
    text = str(value or "").strip()
    if not text:
        raise HTTPException(status_code=400, detail=f"{label} wajib diisi")
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"{label} harus berformat YYYY-MM-DD") from exc


async def record_stack_history(action: str, allocation: dict, operator: str) -> None:
    await db.stack_history.insert_one({
        "id": new_id(), "time": now_iso(), "action": action, "operator": operator,
        "stackCode": allocation.get("stackCode", ""), "allocation": {k: v for k, v in allocation.items() if k != "_id"},
    })


async def _build_allocation(body: StackAllocationBody, allocation_id: Optional[str] = None) -> dict:
    stack_code = body.stackCode.strip().upper()
    if stack_code not in await valid_stack_codes():
        raise HTTPException(status_code=400, detail="Kode tumpukan tidak valid")

    product = await db.products.find_one({"id": body.productId}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")

    per_secondary = float(product.get("secondaryQty", 0) or 0)
    secondary = str(product.get("secondary", "") or "").strip()
    if per_secondary <= 0 or not secondary:
        raise HTTPException(
            status_code=400,
            detail="Atur kemasan sekunder dan isi per kemasan pada master produk terlebih dahulu",
        )

    arrangements = [item.model_dump() for item in body.arrangements] or [{"hamparan": body.length, "kaki": body.width, "height": body.height}]
    secondary_count = sum(item["hamparan"] * item["kaki"] * item["height"] for item in arrangements) + body.extraSecondary
    primary_qty = secondary_count * per_secondary + body.extraPrimary
    query = {"productId": body.productId}
    if allocation_id:
        query["id"] = {"$ne": allocation_id}
    allocated_elsewhere = 0.0
    async for allocation in db.stack_allocations.find(query, {"_id": 0, "primaryQty": 1}):
        allocated_elsewhere += float(allocation.get("primaryQty", 0) or 0)

    stock = float(product.get("stock", 0) or 0)
    if allocated_elsewhere + primary_qty > stock + 1e-9:
        available = max(stock - allocated_elsewhere, 0)
        raise HTTPException(
            status_code=400,
            detail=(
                f"Alokasi melebihi stok {product.get('name', '')}. "
                f"Tersedia untuk ditempatkan {available:g} {product.get('unit', '')}"
            ),
        )

    duplicate_query = {"productId": body.productId, "stackCode": stack_code}
    if allocation_id:
        duplicate_query["id"] = {"$ne": allocation_id}
    if await db.stack_allocations.find_one(duplicate_query):
        raise HTTPException(status_code=409, detail="Produk tersebut sudah tercatat pada tumpukan ini")

    return {
        "productId": product["id"],
        "sku": product.get("sku", ""),
        "productName": product.get("name", ""),
        "unit": product.get("unit", ""),
        "weight": float(product.get("weight", 0) or 0),
        "measureUnit": product.get("measureUnit", "kg") or "kg",
        "secondary": secondary,
        "secondaryQty": per_secondary,
        "stackCode": stack_code,
        "warehouse": stack_code.split("/", 1)[0],
        "zone": re.sub(r"\d", "", stack_code.split("/", 1)[1]),
        "length": body.length,
        "width": body.width,
        "height": body.height,
        "arrangements": arrangements,
        "extraSecondary": body.extraSecondary,
        "extraPrimary": body.extraPrimary,
        "primaryRemainder": body.extraPrimary,
        "secondaryCount": secondary_count,
        "primaryQty": primary_qty,
        "note": body.note.strip(),
        "arrangementAdjusted": False,
        "updatedAt": now_iso(),
    }


async def reconcile_product_allocations(product_id: str) -> None:
    """Pastikan alokasi lokasi tidak pernah lebih besar daripada stok fisik produk."""
    product = await db.products.find_one({"id": product_id}, {"_id": 0, "stock": 1})
    if not product:
        return
    stock = float(product.get("stock", 0) or 0)
    allocations = await db.stack_allocations.find(
        {"productId": product_id}, {"_id": 0}
    ).sort("updatedAt", -1).to_list(5000)
    allocated = sum(float(item.get("primaryQty", 0) or 0) for item in allocations)
    excess = max(allocated - stock, 0)
    if excess <= 1e-9:
        return

    for allocation in allocations:
        if excess <= 1e-9:
            break
        current = float(allocation.get("primaryQty", 0) or 0)
        if excess >= current - 1e-9:
            await db.stack_allocations.delete_one({"id": allocation["id"]})
            await record_stack_history("PENGELUARAN_HABIS", allocation, "Sistem (pengeluaran)")
            excess -= current
            continue

        remaining = current - excess
        per_secondary = float(allocation.get("secondaryQty", 0) or 0)
        await db.stack_allocations.update_one(
            {"id": allocation["id"]},
            {"$set": {
                "primaryQty": remaining,
                "secondaryCount": int(remaining // per_secondary) if per_secondary > 0 else 0,
                "primaryRemainder": remaining % per_secondary if per_secondary > 0 else remaining,
                "length": 0,
                "width": 0,
                "height": 0,
                "arrangementAdjusted": True,
                "updatedAt": now_iso(),
            }},
        )
        await record_stack_history("PENGELUARAN_OTOMATIS", {
            **allocation,
            "primaryQty": remaining,
            "secondaryCount": int(remaining // per_secondary) if per_secondary > 0 else 0,
            "primaryRemainder": remaining % per_secondary if per_secondary > 0 else remaining,
            "length": 0, "width": 0, "height": 0,
            "arrangementAdjusted": True,
        }, "Sistem (pengeluaran)")
        excess = 0


async def decrease_stack_allocation(product_id: str, stack_code: str, qty: float, operator: str) -> None:
    """Kurangi tumpukan asal yang dipilih pada pengeluaran, lalu tandai susunan fisik untuk diperbarui."""
    allocation = await db.stack_allocations.find_one({"productId": product_id, "stackCode": stack_code}, {"_id": 0})
    if not allocation or float(allocation.get("primaryQty", 0) or 0) + 1e-9 < qty:
        raise HTTPException(status_code=400, detail=f"Stok pada tumpukan {stack_code} tidak mencukupi untuk pengeluaran")
    remaining = float(allocation.get("primaryQty", 0) or 0) - qty
    if remaining <= 1e-9:
        await db.stack_allocations.delete_one({"id": allocation["id"]})
        await record_stack_history("PENGELUARAN_HABIS", allocation, operator)
        return
    per_secondary = float(allocation.get("secondaryQty", 0) or 0)
    updated = {**allocation, "primaryQty": remaining, "secondaryCount": int(remaining // per_secondary) if per_secondary > 0 else 0, "primaryRemainder": remaining % per_secondary if per_secondary > 0 else remaining, "length": 0, "width": 0, "height": 0, "arrangementAdjusted": True, "updatedAt": now_iso()}
    await db.stack_allocations.update_one({"id": allocation["id"]}, {"$set": {key: value for key, value in updated.items() if key != "id"}})
    await record_stack_history("PENGELUARAN_OTOMATIS", updated, operator)


async def migrate_default_locations() -> None:
    """Tempatkan stok lama yang sudah mempunyai kode tumpukan default yang valid."""
    legacy = await db.stack_allocations.find({"stackCode": {"$regex": "^MP/"}}, {"_id": 0}).to_list(1000)
    for item in legacy:
        new_code = item["stackCode"].replace("MP/", "MP1/", 1)
        await db.stack_allocations.update_one({"id": item["id"]}, {"$set": {"stackCode": new_code, "warehouse": "MP1"}})
    await db.products.update_many({"location": {"$regex": "^MP/"}}, [{"$set": {"location": {"$replaceOne": {"input": "$location", "find": "MP/", "replacement": "MP1/"}}}}])
    minyak = await db.stack_allocations.find_one({"stackCode": "19/A01", "productName": {"$regex": "MINYAK", "$options": "i"}}, {"_id": 0})
    if minyak and int(minyak.get("length", 0) or 0) == 35 and int(minyak.get("width", 0) or 0) == 19 and int(minyak.get("height", 0) or 0) == 6 and not minyak.get("extraSecondary"):
        product = await db.products.find_one({"id": minyak["productId"]}, {"_id": 0})
        per_secondary = float(minyak.get("secondaryQty", 0) or 0)
        corrected_qty = 4000 * per_secondary
        if product and corrected_qty <= float(product.get("stock", 0) or 0) + 1e-9:
            await db.stack_allocations.update_one({"id": minyak["id"]}, {"$set": {
                "arrangements": [{"hamparan": 35, "kaki": 19, "height": 6}],
                "extraSecondary": 10, "secondaryCount": 4000, "primaryQty": corrected_qty,
            }})
    minyak_2l = await db.stack_allocations.find_one({"stackCode": "19/A02", "productName": {"$regex": r"MINYAK.*2\s*L", "$options": "i"}}, {"_id": 0})
    if minyak_2l and not minyak_2l.get("extraPrimary"):
        product = await db.products.find_one({"id": minyak_2l["productId"]}, {"_id": 0})
        corrected_qty = float(minyak_2l.get("primaryQty", 0) or 0) + 3
        if product and corrected_qty <= float(product.get("stock", 0) or 0) + 1e-9:
            await db.stack_allocations.update_one({"id": minyak_2l["id"]}, {"$set": {
                "extraPrimary": 3, "primaryRemainder": 3, "primaryQty": corrected_qty,
            }})
    products = await db.products.find({"location": {"$in": sorted(await valid_stack_codes())}}, {"_id": 0}).to_list(5000)
    for product in products:
        if not product.get("secondary") or float(product.get("secondaryQty", 0) or 0) <= 0:
            continue
        allocations = await db.stack_allocations.find({"productId": product["id"]}, {"_id": 0, "primaryQty": 1}).to_list(1000)
        unallocated = float(product.get("stock", 0) or 0) - sum(float(x.get("primaryQty", 0) or 0) for x in allocations)
        if unallocated > 1e-9:
            await allocate_stock_to_stack(product, product["location"], unallocated, "Migrasi stok lama")


@router.get("/stack-allocations")
async def list_stack_allocations(user: dict = Depends(get_current_user)):
    return await db.stack_allocations.find({}, {"_id": 0}).sort(
        [("stackCode", 1), ("productName", 1)]
    ).to_list(10000)


@router.post("/stack-allocations")
async def create_stack_allocation(body: StackAllocationBody, user: dict = Depends(require_write)):
    allocation = await _build_allocation(body)
    allocation["id"] = new_id()
    allocation["createdAt"] = now_iso()
    await db.stack_allocations.insert_one(dict(allocation))
    await record_stack_history("DIBUAT", allocation, user.get("name", ""))
    return allocation


@router.put("/stack-allocations/{allocation_id}")
async def update_stack_allocation(
    allocation_id: str,
    body: StackAllocationBody,
    user: dict = Depends(require_write),
):
    existing = await db.stack_allocations.find_one({"id": allocation_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Alokasi tumpukan tidak ditemukan")
    allocation = await _build_allocation(body, allocation_id)
    allocation["createdAt"] = existing.get("createdAt", now_iso())
    result = await db.stack_allocations.update_one({"id": allocation_id}, {"$set": allocation})
    if result.matched_count == 0:
        raise HTTPException(status_code=409, detail="Alokasi berubah. Muat ulang lalu coba lagi")
    updated = {**allocation, "id": allocation_id}
    await record_stack_history("DIUBAH", updated, user.get("name", ""))
    return updated


@router.delete("/stack-allocations/{allocation_id}")
async def delete_stack_allocation(allocation_id: str, user: dict = Depends(require_write)):
    existing = await db.stack_allocations.find_one({"id": allocation_id}, {"_id": 0})
    result = await db.stack_allocations.delete_one({"id": allocation_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Alokasi tumpukan tidak ditemukan")
    await record_stack_history("DIHAPUS", existing, user.get("name", ""))
    return {"ok": True}


@router.get("/stack-treatments")
async def list_stack_treatments(user: dict = Depends(get_current_user)):
    return await db.stack_treatments.find({}, {"_id": 0}).sort("startDate", -1).to_list(10000)


@router.post("/stack-treatments")
async def create_stack_treatment(body: StackTreatmentBody, user: dict = Depends(require_write)):
    warehouse = body.warehouse.strip().upper()
    stack_code = body.stackCode.strip().upper()
    valid_codes = await valid_stack_codes()
    if warehouse not in {code.split("/", 1)[0] for code in valid_codes}:
        raise HTTPException(status_code=400, detail="GBB tidak valid")
    start = _parse_treatment_date(body.startDate, "Tanggal pelaksanaan")
    end = None
    products = []
    standard_days = 0
    if body.type != "SPRAYING":
        end = _parse_treatment_date(body.endDate, "Tanggal buka sungkup")
        if end < start:
            raise HTTPException(status_code=400, detail="Tanggal buka sungkup tidak boleh lebih awal dari tanggal mulai")
        if stack_code not in valid_codes:
            raise HTTPException(status_code=400, detail="Pilih tumpukan untuk fumigasi")
        if stack_code.split("/", 1)[0] != warehouse:
            raise HTTPException(status_code=400, detail="Tumpukan fumigasi harus berada pada GBB yang dipilih")
        allocations = await db.stack_allocations.find({"stackCode": stack_code}, {"_id": 0}).to_list(1000)
        products = [{"productId": x.get("productId"), "name": x.get("productName"), "qty": x.get("primaryQty"), "unit": x.get("unit")} for x in allocations if "BERAS" in str(x.get("productName", "")).upper()]
        if not products:
            raise HTTPException(status_code=400, detail="Fumigasi hanya dapat dicatat pada tumpukan yang berisi beras")
        standard_days = 3 if body.type == "FUMIGASI_SULFUR" else 10
    else:
        stack_code = ""
    normalized_start = start.isoformat()
    normalized_end = end.isoformat() if end else ""
    if await db.stack_treatments.find_one({"type": body.type, "warehouse": warehouse, "stackCode": stack_code, "startDate": normalized_start}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="Catatan pengendalian hama pada tanggal dan lokasi tersebut sudah ada")
    doc = body.model_dump()
    doc.update({
        "id": new_id(),
        "warehouse": warehouse,
        "stackCode": stack_code,
        "scope": "WAREHOUSE" if body.type == "SPRAYING" else "STACK",
        "startDate": normalized_start,
        "endDate": normalized_end,
        "durationDays": (end - start).days if end else 0,
        "standardDurationDays": standard_days,
        "products": products,
        "createdAt": now_iso(),
        "operator": user.get("name", ""),
    })
    await db.stack_treatments.insert_one(dict(doc))
    return doc



def _stack_map_wrap(text: str, font: str, size: float, max_width: float) -> list[str]:
    words = str(text or "").split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if stringWidth(candidate, font, size) <= max_width:
            current = candidate
            continue
        if current:
            lines.append(current)
        current = word
    if current:
        lines.append(current)
    return lines or [""]


def _stack_map_warehouse_config(settings: dict, warehouse: str) -> Optional[dict]:
    warehouses = settings.get("warehouses") or [
        *[
            {
                "code": str(unit),
                "name": f"GBB {unit}",
                "type": "GBB",
                "length": 50,
                "width": 30,
                "zones": [{"code": zone, "count": 4} for zone in ("A", "B", "C")],
                "active": True,
            }
            for unit in range(17, 25)
        ],
        {
            "code": "MP1",
            "name": "MP1",
            "type": "MP",
            "length": 230,
            "width": 30,
            "zones": [{"code": "A", "count": 8}, {"code": "B", "count": 8}],
            "active": True,
        },
    ]
    return next(
        (
            item
            for item in warehouses
            if str(item.get("code", "")).strip().upper() == warehouse and item.get("active", True) is not False
        ),
        None,
    )


def _stack_map_calculation(item: dict) -> str:
    if item.get("arrangementAdjusted"):
        return "Perkalian perlu dihitung ulang"
    arrangements = item.get("arrangements") or [
        {
            "hamparan": item.get("length", 0),
            "kaki": item.get("width", 0),
            "height": item.get("height", 0),
        }
    ]
    parts = []
    for row in arrangements:
        h = int(row.get("hamparan", 0) or 0)
        k = int(row.get("kaki", 0) or 0)
        t = int(row.get("height", 0) or 0)
        parts.append(f"{h} x {k} x {t} = {h * k * t:g}")
    text = " + ".join(parts)
    if float(item.get("extraSecondary", 0) or 0) > 0:
        text += f" + {float(item.get('extraSecondary', 0) or 0):g} {item.get('secondary', 'sekunder')}"
    if float(item.get("extraPrimary", 0) or 0) > 0:
        text += f" + {float(item.get('extraPrimary', 0) or 0):g} {item.get('unit', 'unit')} lepas"
    return text


@router.get("/export/warehouse-stack-map.pdf")
async def export_warehouse_stack_map_pdf(warehouse: str, user: dict = Depends(get_current_user)):
    warehouse = str(warehouse or "").strip().upper()
    settings = await db.settings.find_one({"_id": "app"}, {"_id": 0, "warehouses": 1}) or {}
    config = _stack_map_warehouse_config(settings, warehouse)
    if not config:
        raise HTTPException(status_code=404, detail="Gudang tidak ditemukan")

    zones = [str(zone.get("code", "")).strip().upper() for zone in (config.get("zones") or []) if str(zone.get("code", "")).strip()]
    zone_counts = {
        str(zone.get("code", "")).strip().upper(): int(zone.get("count", 0) or 0)
        for zone in (config.get("zones") or [])
        if str(zone.get("code", "")).strip()
    }
    if not zones or not any(zone_counts.values()):
        raise HTTPException(status_code=400, detail="Konfigurasi tumpukan gudang belum tersedia")

    items = await db.stack_allocations.find(
        {"warehouse": warehouse},
        {"_id": 0},
    ).sort([("stackCode", 1), ("productName", 1)]).to_list(10000)

    grouped: dict[str, list[dict]] = {}
    for item in items:
        code = str(item.get("stackCode", "")).strip().upper()
        grouped.setdefault(code, []).append(item)

    now = operational_now()
    months = [
        "", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
        "Juli", "Agustus", "September", "Oktober", "November", "Desember",
    ]
    pulled = f"{now.day} {months[now.month]} {now.year}"
    printed = f"{now.day} {months[now.month]} {now.year} {now.strftime('%H:%M')} WIB"

    buffer = io.BytesIO()
    page_w, page_h = landscape(A4)
    c = canvas.Canvas(buffer, pagesize=(page_w, page_h))

    # Header
    c.setFillColor(colors.HexColor("#0F172A"))
    c.rect(0, page_h - 31 * mm, page_w, 31 * mm, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 18)
    c.drawCentredString(page_w / 2, page_h - 11.5 * mm, "PETA TUMPUKAN")
    c.setFont("Helvetica-Bold", 10)
    warehouse_name = str(config.get("name") or (f"GBB {warehouse}" if warehouse != "MP1" else "MP1"))
    c.drawCentredString(page_w / 2, page_h - 18 * mm, f"{warehouse_name} - Gudang Sunter Timur I & II")
    c.setFont("Helvetica", 8)
    c.drawCentredString(page_w / 2, page_h - 23.5 * mm, f"Tanggal penarikan data: {pulled}")

    left, right = 12 * mm, page_w - 12 * mm
    front_strip_y = page_h - 40 * mm
    zone_header_y = page_h - 49 * mm
    grid_top = page_h - 57 * mm
    bottom = 24 * mm
    zone_gap = 6 * mm
    usable_w = right - left
    zone_count = max(len(zones), 1)
    col_w = (usable_w - max(zone_count - 1, 0) * zone_gap) / zone_count
    max_rows = max(zone_counts.values())
    row_gap = 2.2 * mm if max_rows > 4 else 3 * mm
    usable_h = grid_top - bottom
    row_h = (usable_h - max(max_rows - 1, 0) * row_gap) / max_rows

    # Door labels get their own strip, so they cannot be covered by zone headers.
    boundary_positions = [
        left + (index + 1) * col_w + (index + 0.5) * zone_gap
        for index in range(zone_count - 1)
    ]
    if warehouse == "MP1" and len(boundary_positions) == 1:
        front_positions = boundary_positions
        back_positions = boundary_positions
    else:
        front_positions = boundary_positions
        back_positions = boundary_positions

    for x in front_positions:
        c.setFillColor(colors.white)
        c.setStrokeColor(colors.HexColor("#CBD5E1"))
        c.roundRect(x - 15 * mm, front_strip_y - 3.3 * mm, 30 * mm, 6.6 * mm, 2 * mm, fill=1, stroke=1)
        c.setFillColor(colors.HexColor("#475569"))
        c.setFont("Helvetica-Bold", 7.2)
        c.drawCentredString(x, front_strip_y - 1.1 * mm, "PINTU DEPAN")

    for zone_index, zone in enumerate(zones):
        x = left + zone_index * (col_w + zone_gap)
        c.setFillColor(colors.HexColor("#DBEAFE"))
        c.roundRect(x, zone_header_y, col_w, 7 * mm, 2 * mm, fill=1, stroke=0)
        c.setFillColor(colors.HexColor("#1D4ED8"))
        c.setFont("Helvetica-Bold", 10)
        c.drawCentredString(x + col_w / 2, zone_header_y + 2.4 * mm, f"TUMPUKAN {zone}")

        for row_index in range(zone_counts.get(zone, 0)):
            code = f"{warehouse}/{zone}{row_index + 1:02d}"
            y = grid_top - (row_index + 1) * row_h - row_index * row_gap
            stack_items = grouped.get(code, [])
            c.setFillColor(colors.HexColor("#F0FDF4") if stack_items else colors.HexColor("#F8FAFC"))
            c.setStrokeColor(colors.HexColor("#86EFAC") if stack_items else colors.HexColor("#CBD5E1"))
            c.roundRect(x, y, col_w, row_h, 2.2 * mm, fill=1, stroke=1)

            c.setFillColor(colors.HexColor("#0F172A"))
            code_font = 8.1 if max_rows > 4 else 9
            c.setFont("Helvetica-Bold", code_font)
            c.drawString(x + 3 * mm, y + row_h - 4.7 * mm, code)

            if not stack_items:
                c.setFillColor(colors.HexColor("#94A3B8"))
                c.setFont("Helvetica-Oblique", 7 if max_rows > 4 else 8)
                c.drawCentredString(x + col_w / 2, y + row_h / 2 - 1 * mm, "KOSONG")
                continue

            visible_items = stack_items[:3]
            hidden_count = max(len(stack_items) - len(visible_items), 0)
            compact = max_rows > 4 or len(visible_items) > 1
            sku_size = 4.4 if compact else 6.3
            name_size = 4.5 if compact else 6.6
            body_size = 4.15 if compact else 5.9
            step = 1.95 * mm if compact else 3.0 * mm
            cursor_y = y + row_h - (7.6 * mm if compact else 9.2 * mm)

            for item_index, item in enumerate(visible_items):
                if item_index:
                    c.setStrokeColor(colors.HexColor("#CBD5E1"))
                    c.line(x + 3 * mm, cursor_y + 0.45 * mm, x + col_w - 3 * mm, cursor_y + 0.45 * mm)
                    cursor_y -= 1.0 * mm

                c.setFillColor(colors.HexColor("#475569"))
                c.setFont("Helvetica-Bold", sku_size)
                c.drawString(x + 3 * mm, cursor_y, f"SKU: {item.get('sku', '-') or '-'}")
                cursor_y -= step

                c.setFillColor(colors.HexColor("#0F172A"))
                c.setFont("Helvetica-Bold", name_size)
                name_lines = _stack_map_wrap(
                    str(item.get("productName", "") or "-"),
                    "Helvetica-Bold",
                    name_size,
                    col_w - 6 * mm,
                )[: 1 if compact else 2]
                for line in name_lines:
                    c.drawString(x + 3 * mm, cursor_y, line)
                    cursor_y -= step

                c.setFillColor(colors.HexColor("#1D4ED8"))
                c.setFont("Helvetica", body_size)
                calculation = _stack_map_calculation(item)
                calc_lines = _stack_map_wrap(calculation, "Helvetica", body_size, col_w - 6 * mm)[:1]
                c.drawString(x + 3 * mm, cursor_y, calc_lines[0] if calc_lines else calculation)
                cursor_y -= step

                qty = float(item.get("primaryQty", 0) or 0)
                unit = str(item.get("unit", "") or "unit")
                c.setFillColor(colors.HexColor("#166534"))
                c.setFont("Helvetica-Bold", body_size)
                c.drawString(x + 3 * mm, cursor_y, f"Total: {qty:g} {unit}")
                cursor_y -= step

                measure_per_unit = float(item.get("weight", 0) or 0)
                measure_total = qty * measure_per_unit
                measure_unit = str(item.get("measureUnit", "kg") or "kg").strip().lower()
                is_oil = "MINYAK" in str(item.get("productName", "")).upper()
                c.setFillColor(colors.HexColor("#334155"))
                c.setFont("Helvetica", body_size)
                if measure_total > 0:
                    if measure_unit in {"liter", "litre", "l"} or is_oil:
                        c.drawString(x + 3 * mm, cursor_y, f"Liter: {measure_total:g} liter")
                    else:
                        c.drawString(x + 3 * mm, cursor_y, f"Berat: {measure_total:g} {item.get('measureUnit', 'kg') or 'kg'}")
                    cursor_y -= step

            if hidden_count > 0:
                c.setFillColor(colors.HexColor("#64748B"))
                c.setFont("Helvetica-Oblique", 4.2 if compact else 5.4)
                c.drawString(x + 3 * mm, max(y + 2.2 * mm, cursor_y), f"+ {hidden_count} komoditi lain")

    back_y = 16 * mm
    for x in back_positions:
        c.setFillColor(colors.white)
        c.setStrokeColor(colors.HexColor("#CBD5E1"))
        c.roundRect(x - 15 * mm, back_y - 3.3 * mm, 30 * mm, 6.6 * mm, 2 * mm, fill=1, stroke=1)
        c.setFillColor(colors.HexColor("#475569"))
        c.setFont("Helvetica-Bold", 7.2)
        c.drawCentredString(x, back_y - 1.1 * mm, "PINTU BELAKANG")

    c.setFillColor(colors.HexColor("#94A3B8"))
    c.setFont("Helvetica", 6.5)
    c.drawRightString(right, 6.5 * mm, f"Dicetak: {printed}")
    c.save()
    buffer.seek(0)

    filename = f"peta_tumpukan_{warehouse}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export/stack-card.xlsx")
async def export_stack_card(stackCode: str, user: dict = Depends(get_current_user)):
    code = stackCode.strip().upper()
    if code not in await valid_stack_codes():
        raise HTTPException(status_code=400, detail="Kode tumpukan tidak valid")
    items = await db.stack_allocations.find({"stackCode": code}, {"_id": 0}).sort("productName", 1).to_list(1000)
    headers = ["KARTU TUMPUKAN", code, "", "", "", "", "", ""]
    rows = [["No", "SKU", "Nama Komoditas", "Susunan P×L×T", "Kemasan Sekunder", "Jumlah Primer", "Satuan", "Kuantum Fisik"]]
    for index, item in enumerate(items, 1):
        parts = [f"{x.get('hamparan', 0)}×{x.get('kaki', 0)}×{x.get('height', 0)}" for x in item.get("arrangements", [])]
        if not parts:
            parts = [f"{item.get('length', 0)}×{item.get('width', 0)}×{item.get('height', 0)}"]
        arrangement = "Perlu dihitung ulang" if item.get("arrangementAdjusted") else " + ".join(parts)
        if not item.get("arrangementAdjusted") and item.get("extraSecondary", 0):
            arrangement += f" + {item.get('extraSecondary')} tambahan"
        loose = f" + {item.get('extraPrimary')} {item.get('unit', '')} lepas" if item.get("extraPrimary", 0) else ""
        rows.append([index, item.get("sku", ""), item.get("productName", ""), arrangement,
                     f"{item.get('secondaryCount', 0)} {item.get('secondary', '')}{loose}", item.get("primaryQty", 0),
                     item.get("unit", ""), f"{float(item.get("primaryQty", 0) or 0) * float(item.get("weight", 0) or 0):g} {item.get("measureUnit", "kg")}"])
    rows.append([])
    rows.append(["RIWAYAT PERUBAHAN SUSUNAN"])
    rows.append(["Waktu", "Aksi", "Produk", "Perkalian", "Jumlah Primer", "Operator"])
    history = await db.stack_history.find({"stackCode": code}, {"_id": 0}).sort("time", -1).to_list(5000)
    for entry in history:
        snap = entry.get("allocation", {})
        calculation = " + ".join(f"{x.get('hamparan')}×{x.get('kaki')}×{x.get('height')}" for x in snap.get("arrangements", []))
        rows.append([entry.get("time", ""), entry.get("action", ""), snap.get("productName", ""), calculation, snap.get("primaryQty", 0), entry.get("operator", "")])
    rows.append([])
    rows.append(["RIWAYAT SPRAYING DAN FUMIGASI"])
    rows.append(["Jenis", "GBB", "Tumpukan", "Mulai", "Selesai/Buka Sungkup", "Komoditas Beras", "Catatan", "Petugas"])
    warehouse = code.split("/", 1)[0]
    treatments = await db.stack_treatments.find({"$or": [{"stackCode": code}, {"type": "SPRAYING", "warehouse": warehouse}]}, {"_id": 0}).sort("startDate", -1).to_list(5000)
    for entry in treatments:
        rows.append([entry.get("type", ""), entry.get("warehouse", ""), entry.get("stackCode", "") or "Semua", entry.get("startDate", ""), entry.get("endDate", ""), ", ".join(x.get("name", "") for x in entry.get("products", [])), entry.get("note", ""), entry.get("operator", "")])
    output = build_xlsx(headers, rows, "Kartu Tumpukan")
    output.seek(0)
    workbook = load_workbook(output)
    sheet = workbook["Kartu Tumpukan"]
    for column in ("C", "D", "F", "G"):
        for cell in sheet[column]:
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical="top",
                wrap_text=True,
            )
    sheet.column_dimensions["C"].width = 48
    sheet.column_dimensions["D"].width = 34
    sheet.column_dimensions["F"].width = 42
    sheet.column_dimensions["G"].width = 42
    for row_index in range(1, sheet.max_row + 1):
        product_text = str(sheet.cell(row=row_index, column=3).value or "")
        if len(product_text) > 45:
            lines = max(2, (len(product_text) + 44) // 45)
            sheet.row_dimensions[row_index].height = min(90, 15 * lines)
    wrapped_output = io.BytesIO()
    workbook.save(wrapped_output)
    wrapped_output.seek(0)
    output = wrapped_output
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="kartu_tumpukan_{code.replace("/", "-")}.xlsx"'})
